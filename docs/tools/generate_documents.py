from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
DOCS_DIR = ROOT / "docs"
REPORT_PATH = DOCS_DIR / "MiniBanking_Projektni_izvestaj.docx"
GUIDE_PATH = DOCS_DIR / "MiniBanking_Priprema_za_odbranu.docx"

# compact_reference_guide preset - exact design tokens
PAGE_WIDTH_IN = 8.5
PAGE_HEIGHT_IN = 11.0
MARGIN_IN = 1.0
HEADER_FOOTER_IN = 0.492
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGIN_TOP_BOTTOM_DXA = 80
CELL_MARGIN_START_END_DXA = 120
LIST_MARKER_DXA = 269  # 0.187 in
LIST_TEXT_DXA = 540    # 0.375 in
LIST_HANGING_DXA = 271

NAVY = RGBColor(11, 37, 69)
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(91, 103, 117)
BODY = RGBColor(24, 34, 46)
WHITE = RGBColor(255, 255, 255)
TABLE_HEADER_FILL = "E8EEF5"
TABLE_ALT_FILL = "F7F9FB"
CALLOUT_FILL = "F4F6F9"
CODE_FILL = "F7F8FA"
BORDER = "B7C5D5"
POSITIVE = RGBColor(31, 58, 95)
CAUTION = RGBColor(122, 90, 0)
RISK = RGBColor(155, 28, 28)


def set_run_font(
    run,
    name: str = "Calibri",
    size: float | None = None,
    color: RGBColor | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_style_language(style, language: str = "sr-Latn-RS") -> None:
    r_pr = style.element.get_or_add_rPr()
    language_element = r_pr.find(qn("w:lang"))
    if language_element is None:
        language_element = OxmlElement("w:lang")
        r_pr.append(language_element)
    language_element.set(qn("w:val"), language)


def configure_styles(document: Document) -> None:
    styles = document.styles

    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = BODY
    set_style_language(normal)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.widow_control = True

    heading_tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for style_name, (size, color, before, after) in heading_tokens.items():
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        set_style_language(style)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.widow_control = True

    custom_styles = {
        "Cover Kicker": (10.5, BLUE, True, False),
        "Cover Title": (29, NAVY, True, False),
        "Cover Subtitle": (14, MUTED, False, False),
        "Metadata": (10.5, BODY, False, False),
        "Table Text": (9.5, BODY, False, False),
        "Code Block": (9.0, BODY, False, False),
        "Question": (11.5, DARK_BLUE, True, False),
    }
    for style_name, (size, color, bold, italic) in custom_styles.items():
        if style_name not in styles:
            style = styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        else:
            style = styles[style_name]
        style.font.name = "Consolas" if style_name == "Code Block" else "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = bold
        style.font.italic = italic
        set_style_language(style)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(4 if style_name != "Cover Title" else 8)
        style.paragraph_format.line_spacing = 1.0 if style_name in {"Code Block", "Cover Title"} else 1.15
        style.paragraph_format.widow_control = True

    styles["Cover Kicker"].paragraph_format.space_after = Pt(6)
    styles["Cover Subtitle"].paragraph_format.space_after = Pt(18)
    styles["Metadata"].paragraph_format.space_after = Pt(3)
    styles["Question"].paragraph_format.space_before = Pt(8)
    styles["Question"].paragraph_format.space_after = Pt(3)
    styles["Question"].paragraph_format.keep_with_next = True


def configure_page(document: Document, running_title: str) -> None:
    section = document.sections[0]
    section.page_width = Inches(PAGE_WIDTH_IN)
    section.page_height = Inches(PAGE_HEIGHT_IN)
    section.top_margin = Inches(MARGIN_IN)
    section.right_margin = Inches(MARGIN_IN)
    section.bottom_margin = Inches(MARGIN_IN)
    section.left_margin = Inches(MARGIN_IN)
    section.header_distance = Inches(HEADER_FOOTER_IN)
    section.footer_distance = Inches(HEADER_FOOTER_IN)

    header = section.header
    header_p = header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_p.paragraph_format.space_after = Pt(0)
    run = header_p.add_run(running_title)
    set_run_font(run, size=8.5, color=MUTED, bold=True)

    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_p.paragraph_format.space_before = Pt(0)
    footer_p.paragraph_format.space_after = Pt(0)
    run = footer_p.add_run("MiniBanking | 2025/2026 | Strana ")
    set_run_font(run, size=8.5, color=MUTED)
    add_field(footer_p, "PAGE")


def add_field(paragraph, field_code: str) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = f" {field_code} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    display = OxmlElement("w:t")
    display.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    set_run_font(run, size=8.5, color=MUTED)
    run._r.extend([begin, instruction, separate, display, end])


def add_custom_numbering(document: Document, number_format: str) -> int:
    numbering = document.part.numbering_part.element
    existing_abstract = [
        int(element.get(qn("w:abstractNumId")))
        for element in numbering.findall(qn("w:abstractNum"))
    ]
    abstract_id = max(existing_abstract, default=-1) + 1
    existing_nums = [
        int(element.get(qn("w:numId")))
        for element in numbering.findall(qn("w:num"))
    ]
    num_id = max(existing_nums, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi_level = OxmlElement("w:multiLevelType")
    multi_level.set(qn("w:val"), "singleLevel")
    abstract.append(multi_level)

    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), number_format)
    level.append(fmt)
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "\u2022" if number_format == "bullet" else "%1.")
    level.append(level_text)
    suffix = OxmlElement("w:suff")
    suffix.set(qn("w:val"), "tab")
    level.append(suffix)
    justification = OxmlElement("w:lvlJc")
    justification.set(qn("w:val"), "left")
    level.append(justification)

    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), str(LIST_TEXT_DXA))
    tabs.append(tab)
    p_pr.append(tabs)
    indentation = OxmlElement("w:ind")
    indentation.set(qn("w:left"), str(LIST_TEXT_DXA))
    indentation.set(qn("w:hanging"), str(LIST_HANGING_DXA))
    p_pr.append(indentation)
    level.append(p_pr)

    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    r_pr.append(fonts)
    level.append(r_pr)
    abstract.append(level)
    first_num = numbering.find(qn("w:num"))
    if first_num is None:
        numbering.append(abstract)
    else:
        numbering.insert(numbering.index(first_num), abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_element = OxmlElement("w:numId")
    num_id_element.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_element])
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25


def restart_numbering(document: Document, base_num_id: int) -> int:
    numbering = document.part.numbering_part.element
    base_num = next(
        element
        for element in numbering.findall(qn("w:num"))
        if int(element.get(qn("w:numId"))) == base_num_id
    )
    abstract_id = base_num.find(qn("w:abstractNumId")).get(qn("w:val"))
    existing_ids = [
        int(element.get(qn("w:numId")))
        for element in numbering.findall(qn("w:num"))
    ]
    new_num_id = max(existing_ids, default=0) + 1

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(new_num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), abstract_id)
    num.append(abstract_ref)
    level_override = OxmlElement("w:lvlOverride")
    level_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    level_override.append(start_override)
    num.append(level_override)
    numbering.append(num)
    return new_num_id


def add_bullets(document: Document, items: Iterable[str], bullet_num_id: int) -> None:
    for item in items:
        paragraph = document.add_paragraph()
        apply_numbering(paragraph, bullet_num_id)
        paragraph.add_run(item)


def add_numbered(document: Document, items: Iterable[str], decimal_num_id: int) -> None:
    list_num_id = restart_numbering(document, decimal_num_id)
    for item in items:
        paragraph = document.add_paragraph()
        apply_numbering(paragraph, list_num_id)
        paragraph.add_run(item)


def set_cell_margins(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (
        ("top", CELL_MARGIN_TOP_BOTTOM_DXA),
        ("bottom", CELL_MARGIN_TOP_BOTTOM_DXA),
        ("start", CELL_MARGIN_START_END_DXA),
        ("end", CELL_MARGIN_START_END_DXA),
    ):
        element = tc_mar.find(qn(f"w:{tag}"))
        if element is None:
            element = OxmlElement(f"w:{tag}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_cell_fill(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def configure_table(table, widths_dxa: Sequence[int], header: bool = True) -> None:
    if sum(widths_dxa) != CONTENT_WIDTH_DXA:
        raise ValueError(f"Table columns must total {CONTENT_WIDTH_DXA}: {widths_dxa}")
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    width = tbl_pr.first_child_found_in("w:tblW")
    if width is None:
        width = OxmlElement("w:tblW")
        tbl_pr.append(width)
    width.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    width.set(qn("w:type"), "dxa")

    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    indent = tbl_pr.first_child_found_in("w:tblInd")
    if indent is None:
        indent = OxmlElement("w:tblInd")
        tbl_pr.append(indent)
    indent.set(qn("w:w"), str(TABLE_INDENT_DXA))
    indent.set(qn("w:type"), "dxa")

    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), BORDER)

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for column_width in widths_dxa:
        grid_column = OxmlElement("w:gridCol")
        grid_column.set(qn("w:w"), str(column_width))
        grid.append(grid_column)

    for row_index, row in enumerate(table.rows):
        row_properties = row._tr.get_or_add_trPr()
        cannot_split = OxmlElement("w:cantSplit")
        cannot_split.set(qn("w:val"), "true")
        row_properties.append(cannot_split)
        if header and row_index == 0:
            set_repeat_table_header(row)
        for column_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            tc_width = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            tc_width.set(qn("w:w"), str(widths_dxa[column_index]))
            tc_width.set(qn("w:type"), "dxa")
            cell.width = Inches(widths_dxa[column_index] / 1440)
            if header and row_index == 0:
                set_cell_fill(cell, TABLE_HEADER_FILL)
            elif row_index % 2 == 0:
                set_cell_fill(cell, TABLE_ALT_FILL)
            for paragraph in cell.paragraphs:
                paragraph.style = "Table Text"
                paragraph.paragraph_format.space_after = Pt(0)
                if header and row_index == 0:
                    paragraph.paragraph_format.keep_with_next = True
                for run in paragraph.runs:
                    set_run_font(run, size=9.5, color=BODY, bold=(header and row_index == 0))


def add_table(
    document: Document,
    headers: Sequence[str],
    rows: Sequence[Sequence[str]],
    widths_dxa: Sequence[int],
) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for row_values in rows:
        row = table.add_row()
        for index, value in enumerate(row_values):
            row.cells[index].text = value
    configure_table(table, widths_dxa, header=True)
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def add_callout(document: Document, label: str, text: str, kind: str = "info") -> None:
    color = {"info": DARK_BLUE, "positive": POSITIVE, "caution": CAUTION, "risk": RISK}[kind]
    table = document.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    label_run = paragraph.add_run(f"{label}: ")
    set_run_font(label_run, size=10.5, color=color, bold=True)
    text_run = paragraph.add_run(text)
    set_run_font(text_run, size=10.5, color=BODY)
    configure_table(table, [CONTENT_WIDTH_DXA], header=False)
    set_cell_fill(cell, CALLOUT_FILL)
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def add_code_block(document: Document, code: str) -> None:
    table = document.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    paragraph = cell.paragraphs[0]
    paragraph.style = "Code Block"
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(code)
    set_run_font(run, name="Consolas", size=9, color=BODY)
    configure_table(table, [CONTENT_WIDTH_DXA], header=False)
    set_cell_fill(cell, CODE_FILL)
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def add_paragraphs(document: Document, paragraphs: Iterable[str]) -> None:
    for text in paragraphs:
        document.add_paragraph(text)


def add_cover(
    document: Document,
    kicker: str,
    title: str,
    subtitle: str,
    metadata: Sequence[tuple[str, str]],
    lead_label: str,
    lead_text: str,
    editorial: bool,
) -> None:
    if editorial:
        kicker_p = document.add_paragraph(style="Cover Kicker")
        kicker_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        kicker_p.paragraph_format.space_before = Pt(112)
    else:
        kicker_p = document.add_paragraph(style="Cover Kicker")
        kicker_p.paragraph_format.space_before = Pt(28)
    kicker_p.add_run(kicker.upper())

    title_p = document.add_paragraph(style="Cover Title")
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER if editorial else WD_ALIGN_PARAGRAPH.LEFT
    title_p.add_run(title)
    subtitle_p = document.add_paragraph(style="Cover Subtitle")
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER if editorial else WD_ALIGN_PARAGRAPH.LEFT
    subtitle_p.add_run(subtitle)

    if editorial:
        spacer = document.add_paragraph()
        spacer.paragraph_format.space_after = Pt(16)
    for label, value in metadata:
        paragraph = document.add_paragraph(style="Metadata")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if editorial else WD_ALIGN_PARAGRAPH.LEFT
        label_run = paragraph.add_run(f"{label}: ")
        set_run_font(label_run, size=10.5, color=MUTED, bold=True)
        value_run = paragraph.add_run(value)
        set_run_font(value_run, size=10.5, color=BODY)

    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(20 if editorial else 12)
    add_callout(document, lead_label, lead_text, "info")
    document.add_page_break()


def set_document_properties(document: Document, title: str, subject: str) -> None:
    properties = document.core_properties
    properties.title = title
    properties.subject = subject
    properties.author = "MiniBanking project"
    properties.keywords = "MiniBanking, Java, Spring Boot, microservices, faculty project"
    properties.comments = "Generated and visually verified as part of the MiniBanking project."


def new_document(running_title: str, title: str, subject: str) -> tuple[Document, int, int]:
    document = Document()
    configure_styles(document)
    configure_page(document, running_title)
    set_document_properties(document, title, subject)
    bullet_num_id = add_custom_numbering(document, "bullet")
    decimal_num_id = add_custom_numbering(document, "decimal")
    return document, bullet_num_id, decimal_num_id


def build_report() -> None:
    document, bullets, numbers = new_document(
        "MINIBANKING | PROJEKTNI IZVEŠTAJ",
        "MiniBanking - projektni izveštaj",
        "Kratak izveštaj o arhitekturi i realizaciji distribuirane bankarske aplikacije",
    )
    add_cover(
        document,
        "Projektovanje distribuiranih sistema",
        "MiniBanking",
        "Projektni izveštaj o mikroservisnoj bankarskoj platformi",
        [
            ("Akademska godina", "2025/2026"),
            ("Implementacija", "Java 21, Spring Boot 4.1.0 i Spring Cloud 2025.1.2"),
            ("Repozitorijum", "Mini-Banking"),
        ],
        "Sažetak",
        "Sistem realizuje obavezne zahteve zadatka, sva četiri opciona bloka i bonus funkcionalnosti: distribuirano praćenje, metrike, skaliranje, WebSocket obaveštenja, profile, globalnu obradu grešaka i Testcontainers E2E test.",
        editorial=False,
    )

    document.add_heading("1. Cilj i obim projekta", level=1)
    add_paragraphs(document, [
        "MiniBanking je distribuirana aplikacija koja demonstrira kako se poslovni domen banke deli na nezavisne servise, kako se oni pronalaze i komuniciraju bez fiksnih adresa i kako sistem nastavlja kontrolisano da radi kada deo zavisnosti nije dostupan. Implementacija nije zamišljena kao produkciona banka, već kao tehnički kompletna nastavna platforma sa jasno vidljivim obrascima distribuiranih sistema.",
        "Rešenje obuhvata devet Spring Boot aplikacija. Tri infrastrukturne aplikacije obezbeđuju service discovery, centralnu konfiguraciju i jedinstvenu ulaznu tačku. Pet domenskih servisa upravlja klijentima, računima, transakcijama, karticama i obaveštenjima, dok Overview servis agregira podatke iz više izvora. Svaki servis ima sopstvenu odgovornost, konfiguraciju i granicu podataka.",
    ])
    add_callout(
        document,
        "Rezultat verifikacije",
        "Kompletan Maven reactor sadrži 22 unit, integration i Testcontainers E2E testa. Završno lokalno izvršavanje prošlo je sa 0 failure-a, 0 error-a i 0 preskočenih testova.",
        "positive",
    )

    document.add_heading("2. Logička arhitektura", level=1)
    add_code_block(document, """Klijent -> Keycloak -> API Gateway -> Eureka LoadBalancer -> poslovni servis
                                      |-> customer-service
                                      |-> account-service
                                      |-> transaction-service -> account-service
                                      |-> card-service --------> account-service
                                      |-> overview-service ----> customer/account/card/transaction

transaction-service -> RabbitMQ -> notification-service -> WebSocket klijent
svi servisi -> Config Server / Actuator / Zipkin / Prometheus -> Grafana""")
    add_table(
        document,
        ["Komponenta", "Uloga u sistemu"],
        [
            ("Discovery Server", "Eureka registar čuva aktivne instance pod logičkim imenom servisa."),
            ("Config Server", "Isporučuje zajedničku i servisnu YAML konfiguraciju, uključujući docker i production profile."),
            ("API Gateway", "Rutira javne zahteve preko lb:// URI-ja, validira JWT i sprovodi CUSTOMER/ADMIN pravila."),
            ("Customer", "CRUD profila klijenata uz validaciju i jedinstven email."),
            ("Account", "Računi, stanje, uplate/isplate i atomski, idempotentan interni transfer."),
            ("Transaction", "Audit zapis i orkestracija transfera preko OpenFeign klijenta."),
            ("Card", "Izdavanje i životni ciklus tokenizovane kartice uz proveru računa."),
            ("Notification", "Idempotentna obrada Rabbit događaja, trajna i WebSocket obaveštenja."),
            ("Overview", "Resilient agregacija klijenta, računa, kartica i transakcija."),
        ],
        [2400, 6960],
    )

    document.add_heading("3. Domen i vlasništvo podataka", level=1)
    add_paragraphs(document, [
        "Domenski servisi koriste zasebne H2 in-memory baze. Customer servis je vlasnik identiteta i statusa klijenta; Account servis je vlasnik stanja novca i evidencije izvršenih internih transfera; Transaction servis je vlasnik poslovnog audit zapisa; Card servis čuva samo token, maskirani PAN i poslednje četiri cifre; Notification servis čuva obaveštenja i jedinstveni ID obrađenog događaja. Nijedan servis ne čita tabelu drugog servisa.",
        "REST sloj koristi request/response record DTO objekte, dok JPA entiteti ostaju interni. Bean Validation odbija neispravne email adrese, brojeve telefona, valute, iznose i dužine polja pre ulaska u poslovnu logiku. Globalni @RestControllerAdvice vraća ujednačen Spring ProblemDetail sa statusom, naslovom, detaljem, URI-jem zahteva, vremenom i mapom validacionih grešaka.",
    ])
    add_bullets(document, [
        "Novčani iznosi su BigDecimal sa najviše dve decimale, čime se izbegava binarna greška tipa double.",
        "H2 radi u PostgreSQL compatibility režimu, a Hibernate create-drop čini svako pokretanje ponovljivim za demonstraciju.",
        "UUID identifikatori ne zavise od jedne centralne sekvence i pogodni su za granice servisa.",
        "Završena transakcija je neizmenjiv audit u smislu brisanja; brisanje je dozvoljeno samo za neuspešan zapis.",
    ], bullets)

    document.add_heading("4. Ključni poslovni tokovi", level=1)
    document.add_heading("4.1 Transfer novca", level=2)
    add_numbered(document, [
        "Klijent šalje POST /api/transactions/transfers kroz Gateway sa JWT tokenom i idempotencyKey vrednošću.",
        "Transaction servis proverava da li isti ključ već postoji. Identičan ponovljen zahtev vraća isti rezultat; drugačiji sadržaj sa istim ključem vraća konflikt.",
        "OpenFeign klijent pronalazi account-service preko Eureka registra i šalje interni transfer sa ID-jem novog transakcionog zapisa.",
        "Account servis pesimistički zaključava oba računa u stabilnom UUID redosledu, proverava statuse, valutu, limit i raspoloživo stanje, a zatim debit i credit izvršava u jednoj lokalnoj DB transakciji.",
        "Transaction servis zaključuje audit. Tek posle uspešnog DB commita objavljuje dva stabilno identifikovana RabbitMQ događaja, po jedan za pošiljaoca i primaoca.",
        "Notification consumer čuva svaki eventId najviše jednom i novo obaveštenje prosleđuje na /topic/notifications/{customerId}.",
    ], numbers)

    document.add_heading("4.2 Izdavanje kartice i agregirani pregled", level=2)
    add_paragraphs(document, [
        "Card servis pre izdavanja kartice preko Feign-a proverava da račun postoji, da pripada navedenom klijentu i da je aktivan. Ne generiše niti vraća stvarni PAN/CVV; demonstracioni model čuva nasumičan token, masku i poslednje četiri cifre.",
        "Overview servis predstavlja composition pattern. Customer pregled spaja profil, sve račune, kartice i istoriju transakcija. Account pregled spaja jedan račun, njegove kartice i transakcije. Svaki udaljeni poziv ima zaseban Retry i Circuit Breaker, pa nedostupan servis daje praznu sekciju i zapis u degradedServices umesto rušenja kompletnog odgovora.",
    ])

    document.add_heading("5. Distribuirani obrasci", level=1)
    add_table(
        document,
        ["Obrazac", "Realizacija", "Dobit"],
        [
            ("Service discovery", "Eureka registracija i lookup po spring.application.name", "Nema fiksnih hostova i portova između servisa."),
            ("Client-side load balancing", "Spring Cloud LoadBalancer bira Overview instancu", "Horizontalno skaliranje bez promene klijenta."),
            ("Declarative HTTP", "@FeignClient(name = \"account-service\") i drugi klijenti", "Interfejs opisuje ugovor, infrastruktura pravi HTTP poziv."),
            ("Circuit Breaker", "CLOSED/OPEN/HALF_OPEN sa failure threshold 50%", "Sprečava lavinu sporih poziva ka neispravnoj zavisnosti."),
            ("Retry", "3 pokušaja, 250 ms, eksponencijalni backoff x2", "Apsorbuje kratkotrajne mrežne greške."),
            ("Idempotency", "Tri nivoa: API ključ, transfer record i eventId", "Retry ne skida novac i ne pravi obaveštenje dva puta."),
            ("Async event", "Topic exchange, durable queue, retry i DLQ", "Transaction ne čeka isporuku korisničkog obaveštenja."),
        ],
        [1900, 3860, 3600],
    )

    document.add_heading("6. Bezbednost i operativna vidljivost", level=1)
    add_paragraphs(document, [
        "Keycloak je authorization server i izdaje potpisan JWT za realm minibanking. Gateway je OAuth2 Resource Server: proverava potpis preko JWK skupa, issuer claim i rok važenja, zatim realm role mapira u Spring Security authorities. CUSTOMER i ADMIN mogu da pristupe /api/**, ali DELETE zahteva ADMIN. Neautentifikovan zahtev dobija 401, a autentifikovan korisnik bez dovoljne uloge 403.",
        "Actuator na svim aplikacijama izlaže health, info, metrics i prometheus. Micrometer dodaje application tag i custom brojače za ishod transfera i Rabbit događaja. Prometheus periodično scrape-uje metrike, Grafana automatski učitava dashboard, a Zipkin prikazuje trace/span lanac i correlation identifikatore u logovima. Time su pokrivena tri stuba: metrike, distribuirani tragovi i korelisani logovi.",
    ])
    add_callout(
        document,
        "Bezbednosna granica",
        "Direktni portovi servisa izloženi su samo radi Swagger/H2 demonstracije. U produkciji bi bili dostupni isključivo na internoj mreži, a Gateway bi ostao jedina javna ulazna tačka.",
        "caution",
    )

    document.add_heading("7. Deploy, profili i testiranje", level=1)
    add_paragraphs(document, [
        "Svaka Java aplikacija ima multi-stage Dockerfile. Build faza koristi Temurin 21 JDK i Maven cache, a runtime faza manji Temurin 21 JRE image i neprivilegovanog spring korisnika. HEALTHCHECK poziva Actuator. Compose definiše redosled zavisnosti, centralne promenljive, portove i volume-e. Osnovni profil pokreće jezgro; scale dodaje drugu Overview instancu, a observability dodaje Prometheus i Grafanu. Heap je ograničen da bi ceo sistem stabilno radio na studentskom računaru.",
        "Testovi pokrivaju kontekst infrastrukturnih aplikacija, validaciju i CRUD, konkurentno važna pravila računa, idempotency transfera, Feign orkestraciju, bezbednosno mapiranje JWT uloga, fallback agregaciju i event publisher. RabbitNotificationEndToEndTests preko Testcontainers-a pokreće pravi RabbitMQ, šalje duplikat događaja i dokazuje da nastaju jedan H2 zapis i jedna WebSocket poruka.",
    ])

    document.add_heading("8. Matrica zahteva", level=1)
    add_table(
        document,
        ["Kategorija", "Zahtev", "Status", "Dokaz u projektu"],
        [
            ("Obavezno", "Eureka, Gateway YAML lb://, LoadBalancer", "Ispunjeno", "discovery-server, api-gateway i dve Overview instance"),
            ("Obavezno", "Feign, CB + Retry + fallback", "Ispunjeno", "servisni klijenti i centralni overview-service.yml pragovi"),
            ("Obavezno", "CRUD, Validation, JPA, H2", "Ispunjeno", "pet domenskih servisa, DTO validacija i zasebne baze"),
            ("Obavezno", "2+ multi-service operacije i agregacija", "Ispunjeno", "transfer, kartica, customer/account overview"),
            ("Obavezno", "Actuator i Swagger", "Ispunjeno", "health/info/metrics i Swagger UI svakog poslovnog servisa"),
            ("Opcije", "Config, RabbitMQ, Docker, JWT", "4/4", "sva četiri opciona bloka implementirana"),
            ("Bonus", "Tracing, metrike, skaliranje, WS, profili", "Ispunjeno", "Zipkin, Prometheus/Grafana, scale, STOMP i docker/production profili"),
            ("Bonus", "Global errors i Testcontainers", "Ispunjeno", "ProblemDetail advice i pravi RabbitMQ E2E test"),
        ],
        [1320, 2820, 1320, 3900],
    )

    document.add_heading("9. Ograničenja i dalji razvoj", level=1)
    add_paragraphs(document, [
        "H2 memorijske baze, demo nalozi i javno izloženi direktni portovi prikladni su za fakultetsku demonstraciju, ali ne za produkciju. Produkciona verzija bi koristila zaseban PostgreSQL po servisu, Flyway migracije, tajne iz vault-a, TLS/mTLS, authorization code flow sa PKCE, granularno vlasništvo nad resursima i zatvorenu servisnu mrežu.",
        "Objavljivanje Rabbit događaja posle commita sprečava slanje događaja za rollback-ovanu transakciju, ali između DB commita i uspešne objave i dalje postoji mali failure window. Transactional outbox i publisher worker uklonili bi taj dual-write rizik. Za finansijski sistem dodatno bi bili potrebni immutable ledger sa dvostrukim knjiženjem, saglasnost salda, sagas/kompenzacije, rate limiting, revizija pristupa i produkcioni SLO/alarmi.",
        "U okviru zadatka izabrano rešenje daje pregledan kompromis: dovoljno je malo da se može pokrenuti i objasniti lokalno, a dovoljno kompletno da demonstrira stvarne probleme i obrasce distribuiranih sistema.",
    ])

    document.save(REPORT_PATH)


def build_guide() -> None:
    document, bullets, numbers = new_document(
        "MINIBANKING | PRIPREMA ZA ODBRANU",
        "MiniBanking - detaljna priprema za odbranu",
        "Tehnologije, implementacija, demonstracioni tok i pitanja za usmenu odbranu",
    )
    add_cover(
        document,
        "Praktični tehnički vodič",
        "MiniBanking",
        "Detaljna priprema za usmenu odbranu i demonstraciju projekta",
        [
            ("Akademska godina", "2025/2026"),
            ("Predviđeno izlaganje", "15-20 minuta + pitanja"),
            ("Osnova", "Java 21, Spring Boot, Spring Cloud i Docker Compose"),
        ],
        "Cilj dokumenta",
        "Da možeš samostalno da pokreneš sistem, objasniš svaku korišćenu tehnologiju, pratiš jedan zahtev kroz kod i argumentovano odgovoriš zašto je izabrano baš ovo rešenje.",
        editorial=True,
    )

    document.add_heading("Kako da koristiš ovaj vodič", level=1)
    add_paragraphs(document, [
        "Prvi prolaz čitaj od početka do kraja i istovremeno otvaraj navedene klase u IDE-u. Drugi prolaz radi uz pokrenut sistem: svaku operaciju izvrši preko scripts/demo.ps1 i pronađi odgovarajući log, trace, metric i zapis u bazi. Treći prolaz je simulacija odbrane: za 15-20 minuta ispričaj tok iz poglavlja 17, a zatim naglas odgovori na pitanja iz poglavlja 19 bez gledanja u tekst.",
        "Ne uči anotacije napamet kao izolovane definicije. Profesoru je važnije da povežeš problem, mehanizam i posledicu. Na primer: mrežni poziv može da padne; zato Overview koristi Retry za kratku grešku, Circuit Breaker da ne preoptereti neispravan servis i fallback da vrati parcijalan odgovor. Uvek objasni gde se to tačno vidi u projektu.",
    ])
    add_callout(
        document,
        "Formula dobrog odgovora",
        "1) koji problem rešavamo; 2) kako tehnologija radi; 3) gde je primenjena u MiniBanking-u; 4) koji kompromis ili ograničenje ostaje.",
        "positive",
    )

    document.add_heading("Predlog rasporeda odbrane", level=2)
    add_table(
        document,
        ["Vreme", "Tema", "Šta pokazuješ", "Glavna poruka"],
        [
            ("0-2 min", "Problem i arhitektura", "README Mermaid dijagram", "Devet aplikacija i jasne granice odgovornosti."),
            ("2-5 min", "Infrastruktura", "Eureka, Config i Gateway", "Nema fiksnih servisnih adresa; konfiguracija i ulaz su centralizovani."),
            ("5-10 min", "Poslovni tok", "scripts/demo.ps1 i Swagger", "Transfer je atomski, auditovan i idempotentan."),
            ("10-13 min", "Otpornost i async", "fallback + RabbitMQ UI", "Kvar jednog servisa ne ruši ceo pregled; obaveštenja su odvojena."),
            ("13-16 min", "Bezbednost", "JWT, 401/403 i Keycloak", "Gateway proverava identitet i uloge."),
            ("16-18 min", "Observability", "Zipkin, Prometheus, Grafana", "Svaki problem može da se vidi kroz health, metriku i trace."),
            ("18-20 min", "Testovi i zaključak", "Maven rezultat i E2E", "22 testa, uključujući pravi RabbitMQ kontejner."),
        ],
        [1120, 1680, 2920, 3640],
    )

    document.add_heading("Sadržaj", level=1)
    add_numbered(document, [
        "Velika slika: šta je sistem i zašto mikroservisi",
        "Struktura repozitorijuma, Maven i Java 21",
        "Spring Boot: IoC, DI, bean-ovi, starteri i konfiguracija",
        "REST API, DTO, validacija i ProblemDetail",
        "JPA, Hibernate, H2 i transakcione granice",
        "Preciznost novca, zaključavanje i idempotentnost",
        "Domenski servisi pojedinačno",
        "Spring Cloud Config i profili",
        "Eureka, LoadBalancer, Gateway i OpenFeign",
        "Resilience4j Circuit Breaker, Retry i fallback",
        "RabbitMQ, DLQ i idempotentni consumer",
        "OAuth2, OIDC, JWT, Keycloak i RBAC",
        "Docker, Compose i lokalno pokretanje",
        "Actuator, Micrometer, Prometheus, Grafana i Zipkin",
        "WebSocket/STOMP obaveštenja",
        "Test strategija i Testcontainers",
        "Kompletna priča jednog transfera",
        "Praktičan scenario i demonstracija kvarova",
        "Pitanja profesora i model odgovora",
        "Code tour, produkcioni razvoj i rečnik pojmova",
    ], numbers)

    document.add_heading("1. Velika slika: šta je MiniBanking", level=1)
    add_paragraphs(document, [
        "MiniBanking je mikroservisna bankarska platforma za upravljanje klijentima, računima, transferima, karticama i obaveštenjima. Klijent ne zna adresu svakog servisa. On komunicira sa API Gateway-em na portu 8080, dobija identitet od Keycloak-a, a Gateway preko Eureka registra pronalazi odgovarajuću instancu poslovnog servisa.",
        "Mikroservis ovde nije samo odvojeni Java package. Svaki izvršni servis ima sopstveni Spring ApplicationContext, port, proces/kontejner, bazu, konfiguraciju i životni ciklus. To omogućava nezavisno skaliranje i izolaciju kvara, ali uvodi mrežnu nepouzdanost, eventual consistency, složenije testiranje i potrebu za observability-em. Projekat namerno pokazuje i prednosti i cenu distribucije.",
        "Granice su izabrane po poslovnoj odgovornosti. Account servis jedini menja stanje računa. Transaction servis ne menja tuđu tabelu, već traži od Account servisa da izvrši transfer i zatim čuva svoj audit. Notification servis nije deo kritične putanje novca; informaciju dobija asinhrono. Overview nema svoju bazu jer samo komponuje podatke koji pripadaju drugim servisima.",
    ])
    add_callout(
        document,
        "Rečenica za odbranu",
        "Sistem je mikroservisan zato što su granice deploy-a i podataka odvojene, a komunikacija ide kroz mrežne ugovore; samo veliki broj klasa ili modula ne bi bio dovoljan da ga nazovemo distribuiranim.",
        "info",
    )

    document.add_heading("1.1 Komponente i smerovi komunikacije", level=2)
    add_code_block(document, """JAVNI TOK
Korisnik --token--> Keycloak
Korisnik --Bearer JWT--> API Gateway --lb://service-name--> poslovni servis

SINHRONO IZMEĐU SERVISA
transaction-service --Feign--> account-service
card-service --------Feign--> account-service
overview-service ----Feign--> customer/account/card/transaction

ASINHRONO
transaction-service --transfer.completed--> RabbitMQ --> notification-service
notification-service --/topic/notifications/{customerId}--> WebSocket klijent""")
    add_bullets(document, [
        "Sinhroni REST je pogodan kada pozivalac odmah mora da zna rezultat, kao kod debit/credit operacije.",
        "Asinhrona poruka je pogodna kada posledica može da se obradi kasnije, kao korisničko obaveštenje.",
        "Centralna infrastruktura ne preuzima vlasništvo nad poslovnim podacima: Eureka zna lokacije, Config vrednosti, Gateway rute, a Keycloak identitete.",
    ], bullets)

    document.add_heading("2. Struktura repozitorijuma, Maven i Java 21", level=1)
    add_paragraphs(document, [
        "Korenski pom.xml ima packaging pom i navodi devet modula. To je Maven reactor: jednom komandom Maven izračunava redosled, nasleđuje verzije i gradi sve module. Parent je spring-boot-starter-parent 4.1.0, Spring Cloud BOM upravlja kompatibilnim verzijama Cloud biblioteka, a java.version je 21. Maven Wrapper čuva ponovljiv Maven runtime, pa korisnik ne mora globalno da instalira istu Maven verziju.",
        "Dependency management ne dodaje biblioteku automatski; on određuje verziju kada modul tu zavisnost zatraži. Nasuprot tome, dependency u korenskom dependencies delu nasleđuju moduli. U projektu su Prometheus registry i Zipkin starter zajednički jer svaka aplikacija treba observability. Servisni POM-ovi dodaju samo svoje startere, na primer JPA/H2 ili OpenFeign.",
        "Java 21 donosi dugoročno podržanu platformu i koristi se za record DTO objekte, switch/jezičke mogućnosti platforme i savremeni JVM. Record je dobar za request/response model jer je nominalno nepromenljiv i automatski daje konstruktor, accessore, equals/hashCode i toString. JPA entitet nije record jer Hibernate zahteva drugačiji životni ciklus, identitet i praćenje promena.",
    ])
    add_table(
        document,
        ["Putanja", "Sadržaj"],
        [
            ("infrastructure/", "Discovery Server, Config Server i API Gateway"),
            ("services/", "Šest poslovnih/kompozicionih servisa"),
            ("config-repository/", "Centralni application i servisni YAML fajlovi"),
            ("deploy/keycloak/", "Realm, client, role i demo user definicije"),
            ("deploy/monitoring/", "Prometheus scrape i Grafana provisioning/dashboard"),
            ("scripts/demo.ps1", "Ponovljiv end-to-end demonstracioni scenario"),
            ("compose.yml", "Cela platforma, profili, healthcheck-ovi i volume-i"),
        ],
        [2500, 6860],
    )

    document.add_heading("2.1 Najvažnije Maven komande", level=2)
    add_code_block(document, """# ceo reactor i svih 22 testa
.\\mvnw.cmd -B -ntp test

# jedan modul
.\\mvnw.cmd -B -ntp -pl services/account-service test

# modul i sve što mu je potrebno (-am = also make)
.\\mvnw.cmd -B -ntp -pl services/customer-service -am package

# lokalno pokretanje jednog modula
.\\mvnw.cmd -pl infrastructure/discovery-server spring-boot:run""")
    add_bullets(document, [
        "-B uključuje batch režim pogodan za CI; -ntp isključuje transfer progress i čini log čitljivijim.",
        "-pl bira projekat/modul; -am traži da se izgrade i njegove reactor zavisnosti.",
        "package pravi izvršni Spring Boot JAR, dok test staje nakon faze testiranja.",
    ], bullets)

    document.add_heading("3. Spring Boot: osnova svake aplikacije", level=1)
    document.add_heading("3.1 IoC container i dependency injection", level=2)
    add_paragraphs(document, [
        "Spring ApplicationContext je IoC container: on kreira objekte kojima upravlja, povezuje njihove zavisnosti i kontroliše životni ciklus. Inverzija kontrole znači da klasa ne konstruiše sama AccountRepository ili RabbitTemplate; container joj ih predaje. Dependency injection je konkretan mehanizam te predaje.",
        "Projekat koristi constructor injection. Na primer TransactionService u konstruktoru prima BankTransactionRepository, AccountClient, ApplicationEventPublisher i MeterRegistry. Prednosti su eksplicitne zavisnosti, mogućnost pravljenja objekta u testu, nepromenljiva polja i rano otkrivanje ciklične zavisnosti. Kod jednog konstruktora @Autowired nije potreban.",
        "@SpringBootApplication spaja @Configuration, component scanning i auto-configuration. Component scanning pronalazi @Service, @Repository, @Component, @RestController i @Configuration klase ispod osnovnog package-a. Auto-configuration zatim uslovno pravi infrastrukturu na osnovu classpath-a i properties vrednosti: DataSource kada postoje JPA i H2, MVC kada postoji web starter, RabbitTemplate kada postoji AMQP starter i konekcioni properties.",
    ])

    document.add_heading("3.2 Stereotipi i slojevi", level=2)
    add_table(
        document,
        ["Anotacija", "Značenje u projektu"],
        [
            ("@RestController", "HTTP adapter: mapira zahtev na DTO i vraća JSON/HTTP status."),
            ("@Service", "Poslovna pravila i transakcione use-case operacije."),
            ("@Repository / JpaRepository", "Pristup podacima i prevođenje persistence izuzetaka."),
            ("@Configuration", "Eksplicitne bean definicije, npr. exchange, queue ili SecurityFilterChain."),
            ("@Component", "Infrastrukturna komponenta, npr. Rabbit publisher/consumer ili JWT converter."),
            ("@ConfigurationProperties", "Tipizovano vezivanje i validacija grupe YAML vrednosti."),
        ],
        [2500, 6860],
    )
    add_paragraphs(document, [
        "Slojevi nisu samo estetska podela. Controller ne sadrži poslovnu transakciju, servis ne poznaje HTTP detalje, a repository ne odlučuje da li je poslovno dozvoljeno zatvaranje računa. Takva podela omogućava testiranje pravila bez ponavljanja web infrastrukture i sprečava da se JPA entitet pretvori u javni API ugovor.",
    ])

    document.add_heading("3.3 Konfiguracija, properties i profili", level=2)
    add_paragraphs(document, [
        "Spring Environment spaja više izvora konfiguracije: application.yml u JAR-u, vrednosti dobijene od Config Server-a, aktivni profil, environment promenljive i command-line argumente. Viši prioritet može da pregazi niži bez recompilacije. Zato ista slika radi lokalno i u Docker-u: URL Eureka servera je localhost po default-u, a Compose ga menja na discovery-server preko EUREKA_URL.",
        "@ConfigurationProperties je bolji od mnoštva @Value polja kada vrednosti čine celinu. AccountProperties mapira minibanking.account.max-transfer-amount na BigDecimal i validira da je pozitivan. CardProperties mapira maksimalni dnevni limit i broj godina važenja. Ako centralna konfiguracija nije validna, aplikacija fail-fast prekida start umesto da radi sa opasnom vrednošću.",
        "Spring profil bira skup dodatnih properties vrednosti. Compose aktivira docker profil. application-docker.yml beleži runtime profil i sampling, a application-production.yml spušta tracing sampling i ograničava detalje health odgovora. Profil ne znači odvojeni kod; isti artefakt menja spoljašnju konfiguraciju.",
    ])

    document.add_heading("4. REST API, DTO, validacija i greške", level=1)
    document.add_heading("4.1 REST semantika", level=2)
    add_paragraphs(document, [
        "REST resursi imaju imenice u putanji: /api/customers, /api/accounts i /api/transactions. GET čita i mora biti bez sporednih efekata, POST kreira ili pokreće poslovnu komandu, PUT zamenjuje izmenljiva polja, a DELETE uklanja resurs samo kada poslovno pravilo to dozvoljava. Specifične komande, kao deposit, withdrawal, block, activate, retry i read, modelovane su kao podresurs/action endpoint-i jer nisu obična zamena celog stanja.",
        "Kreiranje vraća 201 Created, body kreiranog DTO-a i Location header. Uspešno brisanje vraća 204 No Content. Neispravan JSON ili validacija vraća 400, nepostojeći UUID 404, konflikt sa stanjem ili jedinstvenim ključem 409, nedostajući token 401, a nedovoljna uloga 403. Ovi statusi omogućavaju klijentu da razlikuje vrste problema bez parsiranja proizvoljne poruke.",
    ])
    add_table(
        document,
        ["Primer", "Zašto je tako modelovan"],
        [
            ("POST /api/customers", "Kreira novi resurs i vraća 201 + Location."),
            ("GET /api/accounts?customerId=...", "Filter je query parametar, resurs ostaje kolekcija računa."),
            ("POST /api/accounts/{id}/deposits", "Uplata je komanda sa iznosom, nije proizvoljno postavljanje balance polja."),
            ("POST /api/transactions/{id}/retry", "Dozvoljen je samo prelaz neuspešne transakcije kroz poslovni use case."),
            ("DELETE /api/cards/{id}", "Gateway traži ADMIN, a domen dodatno zahteva blokiranu/isteklu karticu."),
        ],
        [3500, 5860],
    )

    document.add_heading("4.2 DTO i Bean Validation", level=2)
    add_paragraphs(document, [
        "Request DTO je ugovor na ivici sistema. CustomerRequest koristi @NotBlank, @Size, @Email i @Pattern; CreateTransferRequest zahteva UUID polja, minimalan iznos 0.01 i maksimalno dve decimale; valuta mora da ima tri slova. @Valid na @RequestBody pokreće validaciju pre poziva servisa. To ne zamenjuje poslovnu validaciju: anotacija može da potvrdi da je amount pozitivan, ali samo AccountService zna da li ima dovoljno sredstava i da li je račun aktivan.",
        "Response DTO eksplicitno bira šta napušta servis. CardResponse izlaže panToken, maskedPan i lastFour, ali nikada realan PAN ili CVV. Entity se ne vraća direktno jer bi promene persistence modela nekontrolisano menjale API, lazy relacije bi mogle da se serijalizuju, a osetljiva polja procure.",
    ])
    add_code_block(document, """POST /api/transactions/transfers
{
  "idempotencyKey": "UUID koji generiše klijent",
  "sourceAccountId": "UUID izvornog računa",
  "destinationAccountId": "UUID ciljnog računa",
  "amount": 125.50,
  "description": "Demo transfer za odbranu"
}""")

    document.add_heading("4.3 Globalni ProblemDetail", level=2)
    add_paragraphs(document, [
        "@RestControllerAdvice je globalni presretač izuzetaka za web sloj jednog servisa. ResourceNotFoundException postaje 404, ConflictException i database unique conflict postaju 409, MethodArgumentNotValidException postaje 400 sa violations mapom, a nečitljiv JSON dobija 400. Spring ProblemDetail standardizuje polja status, title, detail, type i instance, dok projekat dodaje timestamp.",
        "Važna razlika: advice ne guta neočekivanu programersku grešku i ne pretvara svaki problem u 200. HTTP status ostaje semantički tačan, a detalji izuzetka i stack trace ne izlažu se korisniku. U produkciji bi error response nosio i correlation/trace ID da podrška odmah pronađe događaj u logovima.",
    ])

    document.add_heading("5. JPA, Hibernate, H2 i transakcione granice", level=1)
    document.add_heading("5.1 Šta tačno rade JPA i Hibernate", level=2)
    add_paragraphs(document, [
        "JPA je specifikacija objektno-relacionog mapiranja i persistence API-ja; Hibernate je konkretna implementacija. @Entity mapira Java klasu na tabelu, @Id identitet, @Column ograničenja, @Enumerated STRING čitljive enum vrednosti, a @Version optimističku verziju gde je primenjena. Spring Data JPA generiše repository implementaciju iz interfejsa JpaRepository i naziva metoda kao findAllByCustomerIdOrderByCreatedAtDesc.",
        "EntityManager unutar persistence context-a prati managed entitete. Kada servis u @Transactional metodi promeni stanje entiteta, Hibernate dirty checking na flush/commit-u generiše UPDATE bez eksplicitnog save poziva. save se koristi za nove/agregatne objekte, a saveAndFlush u Transaction servisu obezbeđuje da transakcioni zapis dobije identitet pre udaljenog poziva.",
        "open-in-view je isključen. To znači da web serializacija ne drži DB sesiju otvorenom posle servisnog sloja i sprečava skrivene lazy query-je iz controllera. DTO mapiranje se završava unutar transakcione granice, pa broj i mesto upita ostaju kontrolisani.",
    ])

    document.add_heading("5.2 H2 podešavanje", level=2)
    add_paragraphs(document, [
        "Svaki domenski servis ima svoj jdbc:h2:mem naziv, PostgreSQL MODE, DB_CLOSE_DELAY=-1 i DB_CLOSE_ON_EXIT=FALSE. Prva opcija omogućava da memorijska baza živi dok radi JVM i kada nema aktivne konekcije; druga sprečava automatsko gašenje pri JVM shutdown hook detaljima. Hibernate ddl-auto=create-drop generiše šemu pri startu i briše je pri gašenju.",
        "H2 je odličan za brzu, izolovanu demonstraciju i testove, ali nije garancija potpune kompatibilnosti sa PostgreSQL-om. SQL dijalekt, locking i concurrency mogu se razlikovati. Produkcioni korak bio bi zaseban PostgreSQL po servisu, verzionisane Flyway migracije i Testcontainers testovi nad istim tipom baze.",
    ])

    document.add_heading("5.3 @Transactional i šta ona garantuje", level=2)
    add_paragraphs(document, [
        "@Transactional na javnoj servisnoj metodi otvara lokalnu DB transakciju kroz Spring proxy. Ako metoda završi normalno, commit-uje se; za neuhvaćen RuntimeException podrazumevano se radi rollback. readOnly=true saopštava nameru i može da optimizuje flush ponašanje. Izolacija i zaključavanje određuju šta konkurentne transakcije vide.",
        "Ona ne pravi distribuiranu transakciju preko HTTP-a ili RabbitMQ-a. TransactionService i AccountService imaju dve odvojene baze i dva procesa. Account servis može uspešno commit-ovati transfer, a mreža prekinuti odgovor. Zato se koriste idempotentni transferId i zapis rezultata: ponovljeni poziv vraća već izvršenu operaciju bez ponovnog debit-a.",
    ])
    add_callout(
        document,
        "Česta zamka",
        "Spring @Transactional radi preko proxy-ja. Poziv anotirane metode iz iste instance (self-invocation) obično zaobilazi proxy. Zbog toga su resilient udaljeni pozivi izdvojeni u poseban ResilientBankingClients bean, a poslovne transakcije izlažu javne servisne metode.",
        "caution",
    )

    document.add_heading("6. Novac, zaključavanje i idempotentnost", level=1)
    document.add_heading("6.1 Zašto BigDecimal", level=2)
    add_paragraphs(document, [
        "double je binarni floating-point i mnoge decimalne vrednosti, poput 0.1, nema tačnu binarnu reprezentaciju. Sabiranje takvih vrednosti može dati 0.30000000000000004. BigDecimal čuva decimalnu vrednost i skalu, pa je prikladan za novac. DTO ograničava dve decimale, a poređenja koriste compareTo umesto equals, jer 10.0 i 10.00 mogu imati različitu skalu, ali jednaku novčanu vrednost.",
        "Pravi finansijski sistem bi dodatno imao Money value object koji zajedno čuva amount i currency, eksplicitna pravila zaokruživanja i verovatno ledger u kome se saldo izvodi iz knjiženja. Ovde Account entity čuva demonstracioni saldo, a transfer dozvoljava samo istu valutu na obe strane.",
    ])

    document.add_heading("6.2 Pessimistic write lock i redosled", level=2)
    add_paragraphs(document, [
        "AccountRepository.findByIdForUpdate koristi @Lock(PESSIMISTIC_WRITE). Baza zaključava izabrani red do završetka transakcije, pa dve konkurentne isplate ne mogu obe pročitati isto početno stanje i napraviti negativan saldo. Transfer zaključava oba računa.",
        "Ako transfer A->B prvo zaključa A, a transfer B->A prvo B, dve niti mogu da čekaju jedna drugu. Projekat zato sortira oba UUID-a po string vrednosti i uvek zaključava manjeg pa većeg. Poslovna uloga source/destination određuje debit/credit tek nakon zaključavanja. Stabilni globalni redosled uklanja ovu uobičajenu deadlock putanju.",
    ])

    document.add_heading("6.3 Tri sloja idempotentnosti", level=2)
    add_table(
        document,
        ["Sloj", "Ključ", "Efekat ponavljanja"],
        [
            ("Javni transfer API", "idempotencyKey iz zahteva", "Isti sadržaj vraća postojeći TransactionResponse sa idempotentReplay=true."),
            ("Account interni transfer", "transaction ID kao transferId", "Vraća sačuvana završna stanja, bez drugog debit/credit-a."),
            ("Rabbit consumer", "stabilni eventId = UUID(transferId + role)", "Jedinstveni DB zapis sprečava drugo obaveštenje i WebSocket push."),
        ],
        [2350, 2700, 4310],
    )
    add_paragraphs(document, [
        "Idempotentnost ne znači da se zahtev nikada ne šalje dva puta; znači da više istih obrada ima isti poslovni efekat kao jedna. Ako se isti ključ pošalje sa drugačijim računom ili iznosom, servis vraća 409 jer bi tiho prihvatanje sakrilo grešku klijenta.",
    ])

    document.add_heading("7. Domenski servisi", level=1)
    document.add_heading("7.1 customer-service", level=2)
    add_paragraphs(document, [
        "Customer servis je referentni CRUD: kreira, čita po ID-u ili email-u, lista, menja i briše klijente. Email je jedinstven i normalizovan, request ima validaciju, a status opisuje lifecycle. /api/customers/welcome je javni demonstracioni endpoint koji vraća vrednost učitanu sa Config Server-a, pa se centralna konfiguracija može dokazati bez čitanja loga.",
        "Važne klase su CustomerController (HTTP ugovor), CustomerService (pravila i mapiranje), CustomerRepository (persistence), Customer entity i CustomerRequest/Response DTO. GlobalExceptionHandler pokazuje zajednički format greške.",
    ])

    document.add_heading("7.2 account-service", level=2)
    add_paragraphs(document, [
        "Account servis je autoritet za stanje novca. Pored CRUD-a, nudi deposit, withdrawal i interni transfer. Klijent ne može PUT zahtevom direktno da postavi balance; stanje se menja samo kroz kontrolisane poslovne operacije. Račun mora biti ACTIVE, iznos unutar centralno konfigurisanog maksimuma, isplata ne sme preći saldo, a zatvaranje/brisanje je dozvoljeno samo sa nultim saldom.",
        "AccountTransferRecord čuva transferId, oba računa, oba klijenta, iznos, valutu, završna stanja i vreme. Taj zapis je dokaz izvršenja i osnova za idempotentan odgovor pri mrežnom retry-u. Pessimistic lock i stabilan redosled štite konkurentno ažuriranje.",
    ])

    document.add_heading("7.3 transaction-service", level=2)
    add_paragraphs(document, [
        "Transaction servis orkestrira, ali ne knjiži novac. Prvo upisuje PENDING audit, zatim Feign-om šalje transfer Account servisu i na osnovu odgovora prelazi u COMPLETED ili hvata RuntimeException i prelazi u FAILED sa skraćenom porukom. Failed transakcija može da se ponovi istim internim ID-jem; completed se ne može obrisati jer predstavlja audit.",
        "Po uspehu servis objavljuje interni TransferCompletedDomainEvent. TransferEventPublisher ga sluša sa @TransactionalEventListener(AFTER_COMMIT), pravi SOURCE i DESTINATION Rabbit poruku i beleži custom Micrometer metrike: broj completed/failed pokušaja i distribuciju iznosa.",
    ])
    add_callout(
        document,
        "Svesna projektna odluka",
        "Greška Account poziva se čuva kao FAILED poslovni rezultat, pa POST i dalje može vratiti 201 za kreirani audit zapis. Status u response-u govori da novac nije prenet; alternativni API dizajn mogao bi vratiti 502/503 uz isti trajni audit.",
        "info",
    )

    document.add_heading("7.4 card-service", level=2)
    add_paragraphs(document, [
        "Card servis izdaje, čita, menja, blokira, aktivira i briše kartice. Pre izdavanja i ponovne aktivacije Feign pozivom proverava Account servis: račun mora postojati, pripadati istom customerId-u i biti aktivan. Dnevni limit ne sme preći centralno konfigurisan maksimum, a datum isteka se računa iz validityYears vrednosti.",
        "Projekat svesno ne obrađuje pravi kartični PAN. Čuva panToken, maskedPan i lastFour, ne čuva CVV. U produkciji bi token dolazio iz PCI DSS usklađenog vault/payment providera, kriptografski zaštićenog okruženja i strogo auditovanog procesa.",
    ])

    document.add_heading("7.5 notification-service", level=2)
    add_paragraphs(document, [
        "Notification servis ima pun CRUD i komandu mark-read, ali njegova važnija uloga je Rabbit consumer. transfer.completed poruku pretvara u persistentno TRANSFER_COMPLETED obaveštenje. createIfEventIsNew vraća Optional: kada je event nov, consumer povećava processed metric i šalje STOMP poruku; kada je duplikat, povećava duplicate metric bez novog zapisa i push-a.",
        "Ovaj redosled je bitan: prvo trajno čuvanje, pa WebSocket. Ako se browser nije povezao, korisnik obaveštenje i dalje može dobiti preko REST-a. WebSocket je kanal trenutne isporuke, a baza je izvor trajnog stanja.",
    ])

    document.add_heading("7.6 overview-service", level=2)
    add_paragraphs(document, [
        "Overview je stateless composition servis. Customer overview poziva četiri servisa, a account overview tri. ResilientBankingClients obavija svaki Feign poziv zasebnim Retry/Circuit Breaker imenom. ServiceCallResult nosi dostupnost, podatak i grešku, a OverviewService od rezultata gradi response i degradedServices listu.",
        "Stateless dizajn omogućava da se pokrenu overview-1 i overview-2 sa istim spring.application.name. Eureka registruje obe instance, Gateway ruta lb://overview-service koristi LoadBalancer, a /api/overview/instance vraća INSTANCE_ID da se round-robin izbor vidi uživo.",
    ])

    document.add_heading("8. Spring Cloud Config i profili", level=1)
    add_paragraphs(document, [
        "Config Server radi u native profilu i čita lokalni config-repository. Klijent ima spring.config.import=optional:configserver:...; pri startu traži kombinaciju application.yml, {service-name}.yml i aktivnog profila. application.yml nosi zajednički Actuator/tracing/Eureka blok, a account-service.yml, overview-service.yml i drugi nose domenske vrednosti.",
        "Centralizacija znači da prag ili poruka nisu kopirani u više JAR-ova. /customer-service/docker na portu 8888 prikazuje rezultat spajanja konfiguracije za servis i profil. U velikom sistemu Config Server bi čitao Git backend, pristup bi bio zaštićen, osetljive vrednosti šifrovane ili preuzete iz secret manager-a, a refresh kontrolisan.",
    ])
    add_table(
        document,
        ["Fajl", "Primer odgovornosti"],
        [
            ("application.yml", "Zajednički Actuator, tracing, metric tag i Eureka URL."),
            ("application-docker.yml", "Runtime marker i sampling za Compose okruženje."),
            ("application-production.yml", "Niži sampling i oprezniji health detalji."),
            ("account-service.yml", "max-transfer-amount = 250000.00."),
            ("card-service.yml", "max-daily-limit i validity-years."),
            ("overview-service.yml", "CB/Retry pragovi i četiri imenovane resilience instance."),
        ],
        [3000, 6360],
    )
    add_callout(
        document,
        "optional:configserver",
        "optional omogućava da aplikacija u izolovanom testu startuje bez Config Server-a koristeći lokalne/default vrednosti. U strogoj produkciji može se ukloniti optional da nedostupna centralna konfiguracija prekine start.",
        "caution",
    )

    document.add_heading("9. Eureka, LoadBalancer, Gateway i OpenFeign", level=1)
    document.add_heading("9.1 Eureka service discovery", level=2)
    add_paragraphs(document, [
        "Svaka klijentska aplikacija registruje instance-id, host/IP, port, status i metadata pod spring.application.name. Eureka klijent periodično obnavlja lease i preuzima registry cache. Pozivalac traži logičko ime, ne localhost:8082. Kada se instanca ugasi i lease istekne ili deregistracija stigne, više se ne bira za nove pozive.",
        "Discovery Server sam ima register-with-eureka=false i fetch-registry=false zato što je jedini lokalni server, ne poslovni klijent. U visoko dostupnom okruženju postojalo bi više peer Eureka servera. Service discovery rešava lokaciju, ali ne garantuje da je poslovna operacija uspešna; zato i dalje trebaju timeout, retry i circuit breaker.",
    ])

    document.add_heading("9.2 Spring Cloud LoadBalancer", level=2)
    add_paragraphs(document, [
        "LoadBalancer dobija listu instanci od DiscoveryClient-a i bira jednu za konkretan poziv. Gateway pretvara lb://overview-service u stvarni HTTP URI izabrane instance. OpenFeign integracija radi isto za @FeignClient(name=\"account-service\"). Lokalni scale profil pokreće dve Overview instance na host portovima 8086 i 8087, ali obe unutar mreže slušaju port 8086 i dele isti servisni naziv.",
        "Load balancing nije isto što i service discovery. Discovery odgovara 'koje instance postoje', a load balancer 'koju od njih koristi ovaj zahtev'. U produkciji algoritam može uzeti u obzir zone, težine, health ili latency; ovde je cilj vidljiv round-robin demonstracioni efekat.",
    ])

    document.add_heading("9.3 API Gateway", level=2)
    add_paragraphs(document, [
        "Gateway je reaktivna Spring WebFlux aplikacija. YAML definiše Path predicates i lb:// URI-je za customer, account, transaction, card, notification/WebSocket i overview. Default filter dodaje X-MiniBanking-Gateway header. Jedna javna tačka pojednostavljuje CORS, autentikaciju, rutiranje i buduće rate limiting politike.",
        "Gateway nije poslovni servis i ne treba da sadrži transfer pravila. Njegov security filter chain dozvoljava OPTIONS, health/info/prometheus, customer welcome i overview instance; DELETE /api/** traži ADMIN; ostali /api/** i /ws/** traže CUSTOMER ili ADMIN; sve drugo je denyAll.",
    ])

    document.add_heading("9.4 OpenFeign", level=2)
    add_paragraphs(document, [
        "OpenFeign je deklarativni HTTP klijent. Java interfejs navodi @FeignClient name, putanju, metod, parametre i DTO povratnu vrednost, a Spring generiše proxy koji serijalizuje zahtev i deserijalizuje odgovor. Ime se integriše sa Eureka/LoadBalancer-om. Time poslovni servis ne sastavlja URL i ne ponavlja low-level HTTP kod.",
        "Feign ne uklanja distribuciju: poziv i dalje može da bude spor, vrati 404/500, izgubi odgovor ili ima nekompatibilan JSON ugovor. Zbog toga DTO ugovor treba verzionisati/testirati, timeout ograničiti, a retry primeniti samo kada je operacija bezbedna ili idempotentna. Overview GET pozivi su prirodno retry-safe; interni transfer je retry-safe zato što nosi transferId.",
    ])

    document.add_heading("10. Resilience4j Circuit Breaker, Retry i fallback", level=1)
    document.add_heading("10.1 Zašto sva tri mehanizma", level=2)
    add_paragraphs(document, [
        "Retry ponavlja kratkotrajno neuspešan poziv. Circuit Breaker posmatra rezultate i, kada je zavisnost trajno loša, brzo prekida nove pozive. Fallback definiše šta aplikacija vraća kada normalan put nije dostupan. Nijedan mehanizam sam ne rešava sve: retry bez granice može pojačati opterećenje, circuit breaker bez fallback-a i dalje ostavlja grešku klijentu, a fallback bez breaker-a čekao bi timeout za svaki zahtev.",
        "U projektu customerLookup, accountLookup, cardLookup i transactionLookup imaju odvojene instance. To je važno jer kvar Card servisa ne sme otvoriti breaker za Customer servis. Fallback potpis prima iste argumente plus Throwable i vraća ServiceCallResult.unavailable sa bezbednom podrazumevanom vrednošću.",
    ])
    add_code_block(document, """sliding-window-size: 10
minimum-number-of-calls: 5
failure-rate-threshold: 50
slow-call-rate-threshold: 50
slow-call-duration-threshold: 2s
wait-duration-in-open-state: 10s
permitted-number-of-calls-in-half-open-state: 3

retry.max-attempts: 3
retry.wait-duration: 250ms
retry.exponential-backoff-multiplier: 2""")

    document.add_heading("10.2 Stanja Circuit Breaker-a", level=2)
    add_numbered(document, [
        "CLOSED: pozivi prolaze; breaker beleži poslednjih 10 ishoda. Pre minimum 5 poziva nema dovoljno uzorka za otvaranje.",
        "OPEN: kada najmanje 50% poziva padne ili je sporo, novi pozivi se odmah odbijaju i fallback se izvršava bez čekanja udaljenog servisa.",
        "HALF_OPEN: posle 10 sekundi dozvoljava 3 probna poziva. Dovoljan uspeh vraća CLOSED; neuspeh ga vraća u OPEN.",
    ], numbers)
    add_paragraphs(document, [
        "Slow-call prag od dve sekunde je važan jer zavisnost može formalno vraćati 200, ali toliko sporo da iscrpi threadove i korisnički timeout. Overview response eksplicitno nosi degradedServices, pa parcijalan rezultat nije tih ili nevidljiv. Klijent može prikazati dostupne podatke i upozorenje za nedostupnu sekciju.",
    ])

    document.add_heading("11. RabbitMQ, DLQ i idempotentni consumer", level=1)
    document.add_heading("11.1 Zašto asinhrona komunikacija", level=2)
    add_paragraphs(document, [
        "Transfer mora odmah da zna da li je novac pomeren, zato je Transaction -> Account sinhroni poziv. Korisničko obaveštenje ne sme da produži ili obori kritičnu putanju, zato se šalje kao događaj. Transaction servis završava svoj audit i objavljuje poruku; Notification servis je obrađuje svojim tempom. Producer i consumer ne moraju istovremeno biti dostupni jer durable queue zadržava poruku.",
        "Ovo je temporal decoupling, ali ne znači da nema ugovora. Obe strane dele semantiku TransferNotificationEvent JSON-a: eventId, transferId, customerId, customerRole, računi, iznos, valuta i completedAt. Promena event šeme mora biti kompatibilna ili verzionisana.",
    ])

    document.add_heading("11.2 Exchange, routing key, queue i binding", level=2)
    add_table(
        document,
        ["Rabbit element", "Ime", "Funkcija"],
        [
            ("Topic exchange", "minibanking.events", "Prima domenske događaje i rutira po routing key obrascu."),
            ("Routing key", "transfer.completed", "Označava uspešno završen transfer."),
            ("Durable queue", "notification.transfer-events", "Čuva poruke namenjene Notification servisu."),
            ("Direct DLX", "minibanking.dlx", "Prima poruke koje su posle retry-a odbijene."),
            ("DLQ", "notification.transfer-events.dlq", "Izoluje poison message za pregled/reprocessing."),
        ],
        [1900, 2900, 4560],
    )
    add_paragraphs(document, [
        "Exchange ne skladišti poruku kao queue; on odlučuje u koje vezane queue-ove poruka ide. Binding povezuje notification.transfer-events sa minibanking.events za transfer.completed. Topic exchange je izabran jer se platforma lako može proširiti obrascima transfer.*, card.* ili account.*. DLX je direct jer se neuspeh šalje na jednu eksplicitnu routing key vrednost.",
        "JacksonJsonMessageConverter pretvara Java event u JSON i nazad. JSON je čitljiv u RabbitMQ UI-u i manje je vezan za Java serijalizaciju. Producer i consumer ipak moraju uskladiti tipove polja, datum i backward compatibility.",
    ])

    document.add_heading("11.3 Retry, acknowledge i DLQ", level=2)
    add_paragraphs(document, [
        "Simple listener container koristi auto acknowledge: poruka se potvrđuje tek kada listener normalno završi. Retry interceptor pokušava najviše tri puta, sa početnim intervalom 500 ms, multiplierom 2 i maksimumom 2 sekunde. default-requeue-rejected=false sprečava beskonačnu vruću petlju. Kada obrada konačno ne uspe, queue dead-letter argumenti šalju poruku na minibanking.dlx i zatim u DLQ.",
        "DLQ nije automatsko rešenje greške; ona je parking zona. Operater treba da pogleda razlog, popravi podatak ili kod, pa kontrolisano replay-uje poruku. U produkciji bi postojao alarm na rast DLQ-a, metadata o broju pokušaja i administrativni re-drive postupak.",
    ])

    document.add_heading("11.4 AFTER_COMMIT i dual-write ograničenje", level=2)
    add_paragraphs(document, [
        "TransactionService prvo objavi Spring domain event unutar lokalne transakcije. TransferEventPublisher koristi @TransactionalEventListener(phase=AFTER_COMMIT), pa Rabbit poruka ne može otići ako se audit DB transakcija rollback-uje. Za svaki transfer pravi dva event-a, a UUID.nameUUIDFromBytes(transferId + role) daje stabilan ID pri ponavljanju.",
        "Ipak, AFTER_COMMIT nije isto što i atomski commit baze i Rabbit-a. Ako JVM padne posle DB commita, a pre uspešne objave, audit postoji bez poruke. Produkcioni odgovor je transactional outbox: u istoj DB transakciji upisuje se outbox red, a poseban pouzdani publisher ga kasnije šalje i označava. Projekat koristi jednostavniji after-commit obrazac i jasno dokumentuje ovaj kompromis.",
    ])
    add_callout(
        document,
        "Tačna formulacija",
        "RabbitMQ tipično daje at-least-once isporuku, ne magično exactly-once izvršenje. MiniBanking postiže exactly-once poslovni efekat obaveštenja deduplikacijom eventId-a u bazi.",
        "positive",
    )

    document.add_heading("12. OAuth2, OIDC, JWT, Keycloak i RBAC", level=1)
    document.add_heading("12.1 Ko je ko", level=2)
    add_table(
        document,
        ["Uloga", "Komponenta u MiniBanking-u", "Odgovornost"],
        [
            ("Resource Owner", "Demo korisnik customer/admin", "Vlasnik identiteta koji se prijavljuje."),
            ("Authorization Server", "Keycloak", "Autentikuje korisnika i izdaje potpisane tokene."),
            ("Client", "minibanking-cli / demo.ps1", "Traži token i šalje ga API-ju."),
            ("Resource Server", "API Gateway", "Validira access token i štiti rute."),
            ("Protected Resource", "REST/WebSocket rute", "Podaci i operacije dostupni uz odgovarajuću ulogu."),
        ],
        [2050, 2900, 4410],
    )
    add_paragraphs(document, [
        "OAuth 2.0 je okvir za autorizaciju: access token daje ograničen pristup resursu. OpenID Connect dodaje identitetski sloj i standardne informacije o korisniku. U ovoj demonstraciji Keycloak realm minibanking ima public client minibanking-cli, realm role CUSTOMER i ADMIN i dva lokalna korisnika.",
        "scripts/demo.ps1 koristi Resource Owner Password grant radi jednostavne terminalske demonstracije. Taj flow izlaže lozinku client aplikaciji i danas se ne preporučuje za produkcioni browser/mobile sistem. Produkciona varijanta bi koristila Authorization Code + PKCE, preusmerenje na Keycloak login, kratke access tokene i kontrolisan refresh token.",
    ])

    document.add_heading("12.2 JWT nije šifrovan", level=2)
    add_paragraphs(document, [
        "JWT ima tri Base64URL dela odvojena tačkama: header, payload i signature. Header navodi algoritam i key ID. Payload nosi claims kao iss, sub, exp i realm_access.roles. Signature dokazuje da sadržaj nije izmenjen i da ga je potpisao vlasnik privatnog ključa. Payload je lako dekodirati, zato se u njega ne stavljaju lozinke ili tajne.",
        "Gateway preuzima Keycloak JWK skup i bira javni ključ prema kid vrednosti. Proverava kriptografski potpis, issuer http://localhost:9090/realms/minibanking i vreme važenja. KeycloakJwtRoleConverter čita realm_access.roles i pravi ROLE_CUSTOMER/ROLE_ADMIN authorities koje Spring Security koristi u hasRole pravilima.",
    ])
    add_code_block(document, """Authorization: Bearer eyJ...

header.payload.signature
payload (primer claims):
  iss = http://localhost:9090/realms/minibanking
  exp = vreme isteka
  realm_access.roles = [CUSTOMER]""")

    document.add_heading("12.3 401, 403 i security filter chain", level=2)
    add_paragraphs(document, [
        "401 Unauthorized znači da nema prihvatljive autentikacije: token nedostaje, istekao je, issuer je pogrešan ili potpis nije validan. 403 Forbidden znači da je identitet poznat, ali nema potrebnu privilegiju. Zato CUSTOMER koji šalje DELETE dobija 403, dok isti zahtev bez tokena dobija 401.",
        "CSRF je isključen jer je Gateway stateless bearer-token API i ne koristi cookie sesiju za autentikaciju. CORS/OPTIONS preflight je dozvoljen. Query parameter bearer token je omogućen prvenstveno zbog jednostavnog WebSocket handshaka; u produkciji je header ili kratkotrajni WebSocket ticket bezbedniji jer URL može završiti u logu ili istoriji.",
    ])
    add_callout(
        document,
        "Defense in depth",
        "Gateway je trenutna bezbednosna granica, a direktni servisni portovi su izloženi radi fakultetske demonstracije. Produkciona mreža bi ih sakrila i po potrebi dodala mTLS ili token relay/service authorization između servisa.",
        "risk",
    )

    document.add_heading("13. Docker, Compose i lokalno pokretanje", level=1)
    document.add_heading("13.1 Image, container i multi-stage build", level=2)
    add_paragraphs(document, [
        "Docker image je nepromenljiv paket filesystem slojeva i metadata; container je pokrenuta instanca image-a sa sopstvenim procesom i mrežnim namespace-om. Dockerfile svake Java aplikacije ima build i runtime fazu. Build koristi eclipse-temurin:21-jdk-alpine, kopira reactor i Maven Wrapper, koristi BuildKit cache za /root/.m2 i pravi samo ciljni modul sa -pl/-am.",
        "Runtime faza koristi manji eclipse-temurin:21-jre-alpine, kopira samo fat JAR, pravi neprivilegovanog spring korisnika, izlaže port i dodaje HEALTHCHECK na /actuator/health. Time Maven, source kod i JDK compiler ne ulaze u završni runtime image. USER spring smanjuje posledice eventualne ranjivosti u odnosu na root proces.",
    ])

    document.add_heading("13.2 Docker Compose", level=2)
    add_paragraphs(document, [
        "Compose opisuje 15 mogućih servisa: devet Java aplikacija, dve platforme (RabbitMQ i Keycloak), Zipkin i monitoring. Svi su na podrazumevanoj projektnoj mreži i međusobno se adresiraju po service name-u, na primer http://discovery-server:8761. Host port mapping služi pristupu sa Windows-a; container port ostaje servisni port.",
        "depends_on sa condition: service_healthy obezbeđuje da Config čeka Eureka, a poslovni servisi Config. Transaction/Notification čekaju Rabbit health. Gateway čeka infrastrukturu i Keycloak proces. Healthcheck ne dokazuje svaku poslovnu funkciju, ali razlikuje pokrenut proces od aplikacije spremne da služi zahteve.",
        "Named volume čuva Rabbit, Prometheus i Grafana podatke između down/up ciklusa. Bind mount daje Config Server-u YAML repozitorijum, Keycloak-u realm import i monitoring alatima provision fajlove. Environment anchor sprečava ponavljanje Eureka/Config/Zipkin/profile vrednosti.",
    ])
    add_table(
        document,
        ["Compose režim", "Komanda", "Namena"],
        [
            ("Osnovni", "docker compose up --build -d", "Jedna Overview instanca, bez Grafana/Prometheus dodatka."),
            ("Load balancing", "docker compose --profile scale up --build -d", "Dodaje overview-service-2 na host portu 8087."),
            ("Puna odbrana", "docker compose --profile scale --profile observability up --build -d", "Dodaje oba bonus profila."),
            ("Zaustavljanje", "docker compose --profile scale --profile observability down", "Čuva named volume-e."),
        ],
        [1800, 3900, 3660],
    )
    add_paragraphs(document, [
        "JAVA_TOOL_OPTIONS ograničava Java servise na Xms64m/Xmx256m, Keycloak ima 128/384 MB, a Zipkin 64/256 MB. Bez eksplicitnog ograničenja JVM u više kontejnera može planirati veliki procenat iste Docker Desktop memorije. Profil scale i monitoring su opcioni kako bi osnovni sistem bio stabilan i na računaru sa oko 4 GB Docker memorije.",
    ])

    document.add_heading("13.3 Redosled lokalnog starta", level=2)
    add_numbered(document, [
        "Proveri docker version, docker compose version i docker run --rm hello-world. Docker Desktop engine mora biti Running.",
        "Pokreni docker compose up --build -d ili puni profil. Prvi build može trajati jer preuzima Maven i image slojeve.",
        "Prati docker compose ps dok Java servisi ne postanu healthy; za problem koristi docker compose logs -f ime-servisa.",
        "Otvori Eureka, zatim customer welcome i Swagger UI. Tek onda pokreni scripts/demo.ps1.",
        "Po završetku koristi down bez -v. Opcija -v je namerno destruktivna za named volume podatke.",
    ], numbers)

    document.add_heading("14. Actuator, Micrometer, Prometheus, Grafana i Zipkin", level=1)
    document.add_heading("14.1 Health, info i metrics", level=2)
    add_paragraphs(document, [
        "Spring Boot Actuator dodaje operativne endpoint-e. health sabira health contributor-e kao disk, DB, Rabbit, Config ili Circuit Breaker i vraća UP/DOWN; info prikazuje naziv, opis i Maven verziju; metrics omogućava pregled meter-a; prometheus izlaže tekstualni format za scraping. Management endpoint nije poslovni API i u produkciji mora imati kontrolisan pristup.",
        "Liveness odgovara da li proces treba restartovati, readiness da li treba primati saobraćaj. Projekat koristi opšti health za Docker healthcheck. Produkciono orkestrirano okruženje bi razdvojilo probe, jer privremeno nedostupan downstream ne mora značiti da proces treba ubiti.",
    ])

    document.add_heading("14.2 Micrometer i custom metrike", level=2)
    add_paragraphs(document, [
        "Micrometer je facade za metrike, analogan ulozi logging facade-a: aplikacija registruje Counter, Gauge, Timer ili DistributionSummary, a konkretan registry izvozi u Prometheus. Zajednički application tag omogućava filtriranje po servisu.",
        "TransactionService registruje minibanking.transfer.attempts sa outcome completed/failed i minibanking.transfer.amount histogram/distribution summary. Notification consumer registruje minibanking.notification.events sa processed/duplicate. Tag ne sme imati neograničen broj vrednosti: customerId ili transactionId bi napravili high cardinality i opteretili time-series bazu.",
    ])
    add_table(
        document,
        ["Tip metrike", "Kada se koristi", "MiniBanking primer"],
        [
            ("Counter", "Monotono raste za broj događaja.", "Transfer attempts i notification events."),
            ("Gauge", "Trenutna vrednost koja može gore/dole.", "Potencijalno broj PENDING zapisa ili queue depth."),
            ("Timer", "Trajanje i broj operacija.", "Automatske HTTP server/client metrike."),
            ("DistributionSummary", "Distribucija vrednosti koje nisu vreme.", "Iznosi uspešnih transfera."),
        ],
        [1850, 3500, 4010],
    )

    document.add_heading("14.3 Prometheus i Grafana", level=2)
    add_paragraphs(document, [
        "Prometheus radi pull model: po scrape interval-u poziva /actuator/prometheus svake aplikacije i čuva time series sa timestamp-om i labelama. Query je PromQL, na primer sum(rate(minibanking_transfer_attempts_total[5m])) by (outcome). Pull olakšava centralnu kontrolu intervala i otkrivanje nedostupnog target-a.",
        "Grafana ne čuva primarne aplikacione metrike; ona koristi provision-ovani Prometheus datasource i dashboard. Dashboard prikazuje zdravlje servisa, HTTP rate/latency i custom brojače. Za odbranu je dovoljno objasniti put: aplikacija -> Micrometer registry -> Actuator endpoint -> Prometheus scrape -> Grafana panel.",
    ])

    document.add_heading("14.4 Zipkin tracing", level=2)
    add_paragraphs(document, [
        "Trace predstavlja jedan end-to-end zahtev, a span jednu operaciju unutar tog zahteva: Gateway prijem, Feign poziv, HTTP server obradu ili Rabbit publish/consume. Svi span-ovi dele traceId, imaju svoj spanId, roditelja, trajanje, servis i tagove. Propagation header prenosi kontekst kroz HTTP, a tracing instrumentacija ga uključuje u log correlation pattern.",
        "Zipkin collector prima span-ove na /api/v2/spans i UI prikazuje waterfall. Dugačak span otkriva latency, error tag kvar, a prekid stabla problem sa propagacijom. Sampling 1.0 je koristan u demonstraciji jer snima svaki zahtev; u produkciji application-production profil smanjuje verovatnoću da trošak i skladište rastu linearno sa saobraćajem.",
    ])
    add_callout(
        document,
        "Tri signala",
        "Metrika govori da problem postoji, trace gde je vreme ili greška nastala, a log daje detaljan događaj. Nijedan signal sam ne daje celu sliku.",
        "positive",
    )

    document.add_heading("15. WebSocket i STOMP obaveštenja", level=1)
    add_paragraphs(document, [
        "HTTP request/response zahteva da klijent pita server. WebSocket posle HTTP upgrade handshaka održava dvosmernu TCP vezu, pa server može odmah poslati novu poruku. STOMP je jednostavan messaging protokol iznad WebSocket-a sa konceptima CONNECT, SUBSCRIBE, SEND i MESSAGE.",
        "@EnableWebSocketMessageBroker uključuje Spring messaging infrastrukturu. Endpoint za handshake je /ws, application destination prefix /app, a ugrađeni simple broker obrađuje /topic. Notification consumer posle novog DB zapisa poziva SimpMessagingTemplate.convertAndSend('/topic/notifications/' + customerId, response). Browser pretplaćen na tu temu odmah dobija JSON.",
        "Simple broker je dobar za jednu demonstracionu instancu, ali ne deli subscription state između više Notification instanci. Produkcijski scale-out bi koristio broker relay ka RabbitMQ STOMP pluginu ili drugi eksterni pub/sub sloj, autentifikaciju svake pretplate, origin allow-list umesto '*' i strogo pravilo da korisnik može slušati samo svoj customerId.",
    ])

    document.add_heading("16. Test strategija i Testcontainers", level=1)
    document.add_heading("16.1 Šta testiramo", level=2)
    add_paragraphs(document, [
        "Test piramida kombinuje brze unit testove, Spring integration testove i mali broj skupljih E2E infrastrukturnih testova. Unit test izoluje converter/publisher uz mock zavisnosti. Integration test podiže ApplicationContext, pravu servisnu logiku i H2 repository. Testcontainers E2E dodaje pravi RabbitMQ proces.",
        "Kontekst testovi Discovery/Config/Gateway proveravaju da konfiguracija može da se podigne. Customer/Card/Account testovi proveravaju lifecycle, validna i konfliktna pravila. Transaction testovi proveravaju Feign odgovor, idempotent replay i event publisher. Overview testovi proveravaju agregaciju i degradaciju. Security testovi proveravaju mapiranje realm role u ROLE_*.",
    ])
    add_table(
        document,
        ["Nivo", "Prednost", "Ograničenje", "Primer"],
        [
            ("Unit", "Veoma brz i precizno locira pravilo.", "Mock ne dokazuje pravi framework/protokol.", "KeycloakJwtRoleConverterTests."),
            ("Spring integration", "Pravi DI/JPA/H2 i transakcije.", "I dalje može zameniti udaljene zavisnosti.", "AccountServiceIntegrationTests."),
            ("Testcontainers E2E", "Pravi Rabbit protokol, queue i converter.", "Sporiji i zahteva Docker daemon.", "RabbitNotificationEndToEndTests."),
        ],
        [1550, 2700, 2700, 2410],
    )

    document.add_heading("16.2 Rabbit E2E dokaz", level=2)
    add_paragraphs(document, [
        "RabbitNotificationEndToEndTests pokreće rabbitmq:4-management-alpine kroz RabbitMQContainer. Spring properties se dinamički usmeravaju na mapirani host/port kontejnera. Test šalje isti JSON događaj dva puta kroz pravi exchange/queue, čeka asinhronu obradu i proverava da H2 ima tačno jedno obaveštenje i da je SimpMessagingTemplate pozvan jednom.",
        "Ovo istovremeno proverava deklaraciju exchange/queue/binding-a, Jackson JSON converter, @RabbitListener, DB unique/idempotent logiku i WebSocket sporedni efekat. Testcontainers automatski pravi i briše izolovan kontejner, pa test ne zavisi od ručno podešenog lokalnog queue-a.",
    ])
    add_code_block(document, """.\\mvnw.cmd -B -ntp test

Rezultat završne provere:
Tests: 22
Failures: 0
Errors: 0
Skipped: 0
Reactor: 10/10 SUCCESS""")

    document.add_heading("17. Kompletna priča jednog transfera", level=1)
    add_paragraphs(document, [
        "Ovo je najvažnija priča za odbranu. Otvori navedene klase u IDE-u i prati redosled. Ako možeš bez pomoći da objasniš svaki korak i failure scenario, razumeš srž projekta.",
    ])
    add_numbered(document, [
        "Korisnik se autentikuje u Keycloak-u. Token sadrži issuer, rok važenja i CUSTOMER realm role.",
        "PowerShell ili drugi klijent šalje POST /api/transactions/transfers na Gateway sa Authorization: Bearer tokenom.",
        "SecurityWebFilterChain proverava potpis/issuer/exp i KeycloakJwtRoleConverter mapira CUSTOMER u ROLE_CUSTOMER.",
        "Gateway YAML Path predicate bira transaction-service rutu. lb:// URI traži instance iz Eureka registra, a LoadBalancer bira jednu.",
        "TransactionController Bean Validation proverava UUID-e, amount >= 0.01, najviše dve decimale i description dužinu.",
        "TransactionService.createAndExecute traži postojeći BankTransaction po idempotencyKey-u. Ako postoji isti sadržaj, odmah vraća replay; različit sadržaj je 409 konflikt.",
        "Novi BankTransaction se čuva kao PENDING i flush-uje, pa dobija UUID koji postaje interni transferId.",
        "AccountClient Feign proxy traži account-service u Eureka registru i šalje POST /api/accounts/internal/transfers.",
        "AccountService prvo proverava AccountTransferRecord po transferId-u. To pokriva slučaj da je prethodni odgovor izgubljen posle uspešnog commita.",
        "Oba računa se zaključavaju PESSIMISTIC_WRITE u sortiranom UUID redosledu. Proveravaju se različiti ID-evi, limit 250000, ACTIVE status, ista valuta i dovoljan saldo.",
        "Debit izvora, credit odredišta i AccountTransferRecord nastaju u jednoj lokalnoj DB transakciji. Commit oslobađa lock-ove.",
        "Feign odgovor vraća customer ID-eve, valutu, završna stanja i completedAt. Transaction audit prelazi u COMPLETED i beleži ta stanja.",
        "Spring domain event se objavljuje unutar transakcije; TransactionalEventListener ga obrađuje tek AFTER_COMMIT.",
        "TransferEventPublisher pravi stabilan SOURCE i DESTINATION eventId i šalje JSON na minibanking.events sa transfer.completed routing key-em.",
        "Notification listener preuzima svaki događaj, createIfEventIsNew čuva ga najviše jednom, pa SimpMessagingTemplate šalje novu poruku na customer topic.",
        "Korisnik može rezultat videti kroz REST notification kolekciju, WebSocket push, Overview agregaciju, Rabbit UI, Micrometer counter i Zipkin trace.",
    ], numbers)

    document.add_heading("17.1 Šta se događa pri kvaru", level=2)
    add_table(
        document,
        ["Mesto kvara", "Ponašanje sada", "Zaštita / sledeći korak"],
        [
            ("Token nije validan", "Gateway vraća 401; zahtev ne stiže do servisa.", "JWK/issuer/exp validacija."),
            ("Account nije dostupan", "Transaction audit postaje FAILED sa razlogom.", "retry endpoint koristi isti transfer ID; mogao bi se dodati Feign CB."),
            ("Odgovor izgubljen posle Account commita", "Ponovljen interni poziv vraća AccountTransferRecord.", "Idempotentni transferId sprečava dupli debit."),
            ("Transaction DB rollback", "AFTER_COMMIT listener ne šalje Rabbit događaj.", "Nema lažnog obaveštenja za rollback."),
            ("Pad između DB commita i Rabbit publish", "Audit postoji, događaj može izostati.", "Produkcijski transactional outbox."),
            ("Notification privremeno puca", "Listener retry-uje tri puta.", "Posle toga DLQ i alarm/re-drive."),
            ("Ista Rabbit poruka stigne dva puta", "Drugi eventId ne pravi zapis ni WebSocket push.", "DB unique/idempotent consumer."),
            ("Card servis ne radi pri overview-u", "Odgovor sadrži račune i transakcije, cards prazno i degradedServices.", "Retry + Circuit Breaker + fallback."),
        ],
        [2280, 3440, 3640],
    )

    document.add_heading("18. Praktična demonstracija", level=1)
    document.add_heading("18.1 Pre časa", level=2)
    add_bullets(document, [
        "Restartuj računar i proveri da je Docker Desktop engine Running, ne samo da je UI otvoren.",
        "Pokreni puni Compose profil ranije; prvi build ne ostavljaj za trenutak odbrane.",
        "Proveri docker compose ps, Eureka registraciju i Gateway health.",
        "Pokreni .\\mvnw.cmd -B -ntp test i sačuvaj terminal sa BUILD SUCCESS rezultatom.",
        "Otvori unapred tabove: README, Eureka, Customer Swagger, RabbitMQ, Zipkin i Grafana.",
        "Pokreni scripts/demo.ps1 jednom probno, zatim restartuj H2 servise ako želiš potpuno čist demo.",
        "Isključi VPN/proxy ako ometa Docker image ili localhost komunikaciju i priključi laptop na napajanje.",
    ], bullets)

    document.add_heading("18.2 Pokretanje", level=2)
    add_code_block(document, """# iz korena Mini-banking repozitorijuma
docker compose --profile scale --profile observability up --build -d
docker compose ps

# kada su servisi healthy
.\\scripts\\demo.ps1

# log jednog servisa
docker compose logs -f transaction-service

# uredno gašenje, bez brisanja volume-a
docker compose --profile scale --profile observability down""")
    add_paragraphs(document, [
        "Demo skripta čeka Gateway i Keycloak, uzima CUSTOMER token, kreira Anu i Marka, otvara dva RSD računa, izdaje debitnu karticu, prenosi 125.50 RSD i ponavlja isti transfer. Prvi response ima idempotentReplay=false, drugi isti transaction ID i true. Posle kratkog čekanja prikazuje customer overview i Rabbit obaveštenje.",
    ])

    document.add_heading("18.3 Dokaz centralne konfiguracije i load balancing-a", level=2)
    add_code_block(document, """# javno, bez tokena
Invoke-RestMethod http://localhost:8080/api/customers/welcome

# sa scale profilom pozovi više puta
1..6 | ForEach-Object {
  Invoke-RestMethod http://localhost:8080/api/overview/instance
}""")
    add_paragraphs(document, [
        "Welcome response treba da navede config-server kao izvor. Instance pozivi treba da smenjuju overview-1 i overview-2. U Eureka UI pokaži da se pod jednim OVERVIEW-SERVICE imenom vide dve instance: host portovi nisu upisani u Gateway rutu.",
    ])

    document.add_heading("18.4 Kontrolisana demonstracija fallback-a", level=2)
    add_numbered(document, [
        "Iz demo output-a kopiraj customerA.id i pripremi validan Bearer token.",
        "Pozovi /api/overview/customers/{id} i pokaži kompletan response.",
        "Izvrši docker compose stop card-service.",
        "Ponovi overview. Posle retry-a response i dalje stiže, cards je prazno, a degradedServices navodi card-service i uzrok.",
        "Ponovi nekoliko puta da breaker pređe OPEN; objasni da kasniji poziv više ne čeka mrežni timeout.",
        "Izvrši docker compose start card-service, sačekaj registraciju i HALF_OPEN probne pozive, pa pokaži oporavak.",
    ], numbers)
    add_callout(
        document,
        "Ne improvizuj destruktivno",
        "Na odbrani za kvar koristi stop/start jednog tačno imenovanog kontejnera. Ne koristi down -v, ne briši image-e i ne menjaj YAML uživo.",
        "risk",
    )

    document.add_heading("18.5 Bezbednosna demonstracija", level=2)
    add_paragraphs(document, [
        "Prvo pozovi /api/customers bez Authorization header-a i pokaži 401. Zatim uzmi CUSTOMER token i uspešno izvrši GET. Konačno pokušaj DELETE sa CUSTOMER tokenom i pokaži 403. Ako želiš uspešan DELETE, koristi admin/admin123 token i resurs koji poslovna pravila dozvoljavaju da se obriše. Time razdvajaš autentikaciju, autorizaciju i domensko pravilo.",
    ])

    document.add_heading("18.6 Ako nešto ne radi", level=2)
    add_table(
        document,
        ["Simptom", "Najverovatniji uzrok", "Provera"],
        [
            ("docker API pipe nije pronađen", "Docker Desktop daemon nije pokrenut.", "Otvori Desktop, čekaj Running, pokreni docker version."),
            ("Gateway 503", "Ciljni servis još nije u Eureka registry-ju.", "Eureka UI, service log i /actuator/health."),
            ("Gateway 401 sa tokenom", "Token issuer ne odgovara localhost realm-u ili je istekao.", "Dekodiraj iss/exp i uzmi nov token."),
            ("Transfer FAILED", "Account nije dostupan, račun/valuta/stanje/limit nije validan.", "failureReason i account-service log."),
            ("Nema notification zapisa", "Rabbit queue/consumer nije spreman ili je poruka u DLQ.", "Rabbit UI, notification log i DLQ."),
            ("Grafana prazna", "Observability profil nije aktivan ili Prometheus target je DOWN.", "Prometheus /targets i Compose profile."),
        ],
        [2400, 3440, 3520],
    )

    document.add_heading("19. Pitanja profesora i model odgovora", level=1)
    add_paragraphs(document, [
        "Odgovore koristi kao model, ne kao tekst koji moraš doslovno izgovoriti. Dobar odgovor je kraći od jednog minuta, a profesorov potpitanje zatim možeš povezati sa konkretnom klasom ili demonstracijom.",
    ])

    def add_qa(question: str, answer: str) -> None:
        question_paragraph = document.add_paragraph(style="Question")
        question_paragraph.add_run(question)
        answer_paragraph = document.add_paragraph()
        label = answer_paragraph.add_run("Model odgovora: ")
        set_run_font(label, size=11, color=BODY, bold=True)
        answer_paragraph.add_run(answer)

    qa_items = [
        (
            "1. Zašto mikroservisi, a ne jedan Spring Boot monolit?",
            "Granice klijenta, računa, transakcije, kartice i obaveštenja imaju odvojeno vlasništvo podataka i deploy. To demonstrira nezavisno skaliranje, izolaciju kvara i različite komunikacione obrasce. Cena je mrežna nepouzdanost, eventual consistency, više infrastrukture i observability-a. Za mali realni proizvod modularni monolit bi često bio jeftiniji; ovde je cilj upravo učenje distribuiranih obrazaca.",
        ),
        (
            "2. Da li je Maven modul isto što i mikroservis?",
            "Nije. Modul je build-time organizacija. Mikroservis mora biti nezavisno izvršan/deployable proces sa sopstvenim runtime lifecycle-om i granicom podataka. Svaki MiniBanking servis proizvodi svoj JAR, sluša svoj port, registruje se u Eureki i ima svoju H2 bazu.",
        ),
        (
            "3. Koja je razlika između Eureka registra i API Gateway-a?",
            "Eureka čuva lokacije i status instanci; ne predstavlja javni poslovni API. Gateway je ulazna tačka koja prima klijentski zahtev, primenjuje security i route pravila, pa koristi discovery/load balancing da pronađe instancu. Registry odgovara 'gde', Gateway 'kuda i pod kojim pravilima'.",
        ),
        (
            "4. Koja je razlika između service discovery-ja i load balancing-a?",
            "Discovery vraća skup dostupnih instanci za logičko ime. Load balancer iz tog skupa bira jednu instancu za konkretan poziv. U projektu Eureka daje overview-1/overview-2, a Spring Cloud LoadBalancer bira instancu iza lb://overview-service.",
        ),
        (
            "5. Zašto Config Server ako postoji application.yml?",
            "Config Server centralizuje vrednosti koje se menjaju po servisu/okruženju i uklanja ih iz JAR-a. Zajednički Actuator/Eureka/tracing blok se ne kopira, a max transfer i resilience pragovi imaju jedno mesto. Lokalni application.yml i dalje sadrži bootstrap/default vrednosti potrebne da servis zna kako da nađe Config Server.",
        ),
        (
            "6. Kako OpenFeign poziv pronalazi Account servis bez URL-a?",
            "@FeignClient(name='account-service') pravi proxy. Spring Cloud LoadBalancer pita Eureka DiscoveryClient za instance account-service, izabere jednu i sastavi stvarni URI. Interfejs zatim opisuje HTTP metodu, putanju, body i response DTO.",
        ),
        (
            "7. Koja su stanja Circuit Breaker-a?",
            "CLOSED propušta i meri pozive. Kada prag neuspeha/sporih poziva pređe 50% uz najmanje pet poziva, prelazi OPEN i brzo odbija. Posle 10 sekundi HALF_OPEN propušta tri probe; uspeh zatvara, neuspeh ponovo otvara breaker.",
        ),
        (
            "8. Zašto Retry može da bude opasan?",
            "Povećava ukupan broj zahteva baš kada je zavisnost opterećena i može ponoviti ne-idempotentnu operaciju. Zato je broj ograničen, postoji backoff, a koristi se na GET lookup pozivima. Transfer se može bezbedno ponoviti samo zato što transferId ima idempotentni zapis.",
        ),
        (
            "9. Da li fallback sakriva problem?",
            "Može ako vrati lažni normalan podatak. MiniBanking vraća praznu bezbednu sekciju, ali istovremeno popunjava degradedServices sa imenom i razlogom. Klijent i monitoring zato vide degradaciju, dok dostupni delovi odgovora ostaju korisni.",
        ),
        (
            "10. Da li @Transactional pokriva Transaction i Account servis zajedno?",
            "Ne. Spring transakcija je lokalna za jednu bazu i jedan process/proxy kontekst. HTTP poziv prelazi tu granicu. Sistem koristi idempotentne komande i audit, a za složeniji multi-step proces koristio bi saga/kompenzacije, ne jednu skrivenu JPA transakciju.",
        ),
        (
            "11. Kako sprečavate race condition pri isplati?",
            "Account red se čita sa PESSIMISTIC_WRITE lock-om unutar transakcije. Druga izmena istog računa čeka, pa proveru salda i debit izvršavamo nad konzistentnim stanjem. Bez lock-a dve niti bi mogle pročitati isti saldo i obe proći proveru.",
        ),
        (
            "12. Kako sprečavate deadlock kod dva računa?",
            "Svaki transfer sortira oba UUID-a i zaključava ih istim globalnim redosledom, nezavisno od smera transfera. Zato A->B i B->A ne uzimaju prvo različite lock-ove. DB deadlock je i dalje opšti mogući događaj, ali ova tipična putanja je uklonjena.",
        ),
        (
            "13. Zašto BigDecimal, a ne double?",
            "double ne predstavlja tačno većinu decimalnih iznosa i uvodi kumulativnu rounding grešku. BigDecimal čuva decimalnu vrednost/skalu i podržava eksplicitna pravila. Poređenje koristimo preko compareTo, a validacija dozvoljava dve decimale.",
        ),
        (
            "14. Šta je idempotentnost u ovom projektu?",
            "Ponovljena ista operacija daje isti poslovni efekat kao jedno izvršenje. Javni idempotencyKey vraća isti Transaction, interni transferId vraća AccountTransferRecord bez drugog debit-a, a Rabbit eventId sprečava drugo obaveštenje. Isti ključ sa drugačijim sadržajem je konflikt.",
        ),
        (
            "15. Da li RabbitMQ garantuje exactly-once?",
            "Ne računamo na exactly-once delivery. Poruka može stići ponovo zbog ack/reconnect scenarija. Sistem koristi at-least-once pristup i idempotentni consumer sa jedinstvenim eventId-jem, pa postiže exactly-once poslovni efekat za zapis/push.",
        ),
        (
            "16. Zašto događaj šaljete AFTER_COMMIT?",
            "Da consumer ne dobije obaveštenje za DB transakciju koja se kasnije rollback-uje. To ipak ostavlja prozor između commita i Rabbit publish-a; produkcijski transactional outbox bi u istoj transakciji upisao događaj za kasnije pouzdano slanje.",
        ),
        (
            "17. Čemu služi DLQ?",
            "Poruka koja i posle ograničenog retry-a ne može da se obradi ne treba beskonačno da blokira ili kruži. DLQ je izoluje za pregled, alarm i kontrolisan re-drive nakon popravke podataka ili koda.",
        ),
        (
            "18. Koja je razlika topic i direct exchange-a?",
            "Direct zahteva tačno poklapanje routing key-a, topic podržava obrasce sa * i # segmentima. Domenski događaji koriste topic da bi se kasnije pretplatilo na transfer.*; DLX koristi direct jer neuspeh ide na eksplicitnu queue putanju.",
        ),
        (
            "19. Da li je JWT šifrovan?",
            "Ne nužno; naš JWT je potpisan. Payload se može Base64URL dekodirati, pa u njega ne stavljamo tajne. Potpis štiti integritet i autentičnost, a TLS štiti token u transportu.",
        ),
        (
            "20. OAuth2 naspram OpenID Connect-a?",
            "OAuth2 standardizuje delegiranu autorizaciju i access token. OIDC dodaje identitet, ID token, user-info i standardne authentication semantike iznad OAuth2. Keycloak podržava oba; Gateway je resource server za access token.",
        ),
        (
            "21. Razlika 401 i 403?",
            "401 znači da zahtev nema validnu autentikaciju. 403 znači da je korisnik autentifikovan, ali nema dovoljnu ulogu. CUSTOMER DELETE je 403; zahtev bez tokena je 401.",
        ),
        (
            "22. Da li je Gateway single point of failure?",
            "Jedna lokalna instanca jeste. Produkciono bi se pokrenulo više stateless Gateway instanci iza spoljnog L4/L7 load balancer-a, sa zajedničkim discovery/config/security pravilima. Lokalni projekat optimizuje demonstracionu jednostavnost.",
        ),
        (
            "23. Zašto su direktni portovi servisa javni ako Gateway ima security?",
            "Samo radi Swagger/H2 fakultetske demonstracije na localhost-u. To nije production topology. U produkciji se ne bi publish-ovali host portovi, mrežne politike bi dozvolile samo Gateway-u/ovlašćenim servisima, a mogao bi se dodati mTLS ili resource-server zaštita svakog servisa.",
        ),
        (
            "24. Zašto H2 i koje je ograničenje?",
            "Zadatak traži H2, brz je, nema spoljašnju instalaciju i svaki test dobija čistu bazu. Ograničenje je memorijska netrajnost i razlike u SQL/locking ponašanju u odnosu na produkcioni PostgreSQL. Zato bi produkcija koristila zasebne baze, migracije i real-DB Testcontainers testove.",
        ),
        (
            "25. Zašto je open-in-view isključen?",
            "DB transakcija ne treba da curi do JSON serializacije. Isključivanje sprečava skrivene lazy query-je/N+1 iz web sloja i tera nas da učitamo/mapiramo DTO unutar servisa. Time su granice i performanse predvidljiviji.",
        ),
        (
            "26. Šta dobijamo sa ProblemDetail?",
            "Konzistentan, mašinski čitljiv error format sa status/title/detail/type/instance, uz violations i timestamp. Klijent ne mora posebno parsirati poruke svakog endpoint-a, a HTTP status ostaje tačan.",
        ),
        (
            "27. Razlika Docker image i container-a?",
            "Image je nepromenljivi template sa slojevima i konfiguracijom; container je njegova pokrenuta instanca sa procesom, writable slojem i mrežom. Iz istog Overview image-a Compose pravi dve instance sa različitim INSTANCE_ID-em.",
        ),
        (
            "28. Čemu služe multi-stage Dockerfile i non-root USER?",
            "Build alati/JDK/source ostaju u build fazi, a runtime image dobija samo JRE i JAR, pa je manji i ima manju attack surface. Non-root spring korisnik smanjuje privilegije kompromitovanog procesa.",
        ),
        (
            "29. Da li depends_on znači da je servis potpuno spreman?",
            "Samo uz condition: service_healthy čeka definisan healthcheck; service_started čeka proces. Čak ni health ne garantuje svaku poslovnu zavisnost. Zato klijenti moraju podneti kasnu registraciju i privremenu nedostupnost.",
        ),
        (
            "30. Razlika Actuator, Prometheus i Grafana?",
            "Actuator/Micrometer proizvode i izlažu metrike; Prometheus ih periodično preuzima, skladišti i upituje; Grafana ih vizuelizuje kroz dashboard. Grafana nije izvor podataka, a Actuator nije time-series baza.",
        ),
        (
            "31. Šta su trace i span?",
            "Trace je celo putovanje jednog zahteva kroz servise. Span je jedna vremenski ograničena operacija sa parent/child odnosom. traceId povezuje sve span-ove i logove, a spanId identifikuje konkretnu operaciju.",
        ),
        (
            "32. Zašto Testcontainers ako imamo mock test?",
            "Mock proverava našu interakciju sa zamišljenim API-jem, ali ne proverava Rabbit handshake, deklaracije, binding, JSON converter ili realnu asinhronost. Testcontainers pokreće pravi broker u ponovljivom izolovanom okruženju i hvata konfiguracione greške.",
        ),
        (
            "33. Koji servis je najlakše horizontalno skalirati?",
            "Overview je stateless i nema bazu, pa se instance mogu dodavati iza istog service name-a bez migracije session state-a. Domenski servisi mogu da se skaliraju uz zajedničku eksternu bazu i pravilno locking/cache ponašanje; njihove sadašnje in-memory H2 baze nisu zajedničke i zato se ne skaliraju za konzistentan dataset.",
        ),
        (
            "34. Da li Account servis proverava da customerId stvarno postoji?",
            "Trenutna demonstracija čuva customerId kao međuservisnu referencu bez sinhrone provere pri otvaranju računa; demo/orchestrator prvo kreira klijenta. To je dokumentovano ograničenje. Produkciono bi ovlašćeni onboarding use case proverio Customer servis ili bi Account održavao event-driven lokalni read model aktivnih klijenata, uz proceduru za lifecycle događaje.",
        ),
        (
            "35. Šta je saga i gde bi se koristila?",
            "Saga je niz lokalnih transakcija kroz servise sa događajima/komandama i kompenzacionim akcijama umesto globalne ACID transakcije. Bila bi potrebna za složen onboarding ili eksterni transfer sa više koraka. Naš interni transfer je jedna lokalna Account transakcija, pa saga tu nije potrebna.",
        ),
        (
            "36. Kako CAP teorema utiče na projekat?",
            "Pri mrežnoj particiji distribuirani sistem ne može istovremeno garantovati potpunu konzistentnost i dostupnost za isti podatak. Account servis bira strogu lokalnu konzistentnost salda; Overview bira dostupnost parcijalnog prikaza i jasno označava degradaciju. CAP se razmatra po operaciji/podatku, ne jednom etiketom za ceo sistem.",
        ),
    ]
    for question, answer in qa_items:
        add_qa(question, answer)

    document.add_heading("20. Code tour i dalje unapređenje", level=1)
    document.add_heading("20.1 Putanja kroz kod za pet minuta", level=2)
    add_table(
        document,
        ["Tema", "Prvo otvori", "Zatim pokaži"],
        [
            ("Gateway ruta", "api-gateway/application.yml", "SecurityConfiguration i KeycloakJwtRoleConverter"),
            ("Transfer", "TransactionController", "TransactionService -> AccountClient -> AccountService"),
            ("Zaključavanje", "AccountRepository.findByIdForUpdate", "AccountService.transfer i AccountTransferRecord"),
            ("Rabbit", "TransferEventPublisher", "RabbitMessagingConfiguration -> TransferNotificationConsumer"),
            ("Fallback", "ResilientBankingClients", "overview-service.yml -> OverviewService"),
            ("Metrike", "TransactionService constructor", "prometheus.yml i Grafana dashboard JSON"),
            ("E2E test", "RabbitNotificationEndToEndTests", "Testcontainers dependency u notification POM-u"),
        ],
        [1850, 3600, 3910],
    )

    document.add_heading("20.2 Šta bih uradio za produkciju", level=2)
    add_bullets(document, [
        "PostgreSQL po servisu, Flyway migracije, backup/restore vežbe i šifrovanje podataka u mirovanju.",
        "Immutable double-entry ledger, rezervacije sredstava, reconciliation i jasni statusi settlement-a.",
        "Transactional outbox/inbox, schema registry ili verzionisani event ugovori i automatizovan DLQ re-drive.",
        "Authorization Code + PKCE, kratki tokeni, tajne iz Vault/KMS-a, TLS/mTLS, rate limiting i audit pristupa.",
        "Kubernetes ili druga orkestracija, više Gateway/Config/registry instanci, resource requests/limits i autoscaling.",
        "Contract testovi (npr. consumer-driven), PostgreSQL/Rabbit E2E, load/chaos/security testovi i dependency scanning.",
        "OpenTelemetry collector, centralni log sistem, alerting, SLI/SLO i runbook za incident response.",
        "API versioning, pagination, optimistic concurrency/ETag i granularno vlasništvo korisnika nad resursom.",
    ], bullets)

    document.add_heading("20.3 Brzi rečnik", level=2)
    add_table(
        document,
        ["Pojam", "Jednorečenična definicija"],
        [
            ("ACID", "Atomicity, Consistency, Isolation, Durability osobine lokalne transakcije."),
            ("API Gateway", "Jedinstvena kontrolisana ulazna tačka i reverse proxy za API-je."),
            ("Bean", "Objekat čiji životni ciklus i zavisnosti upravlja Spring container."),
            ("Circuit Breaker", "State machine koja brzo prekida pozive ka neispravnoj zavisnosti."),
            ("DTO", "Objekat javnog ugovora za prenos podataka, odvojen od persistence entiteta."),
            ("DLQ", "Queue za poruke koje nisu uspešno obrađene posle retry politike."),
            ("Eventual consistency", "Različiti servisi ne moraju trenutno imati isti pogled, ali konvergiraju."),
            ("Feign", "Deklarativni HTTP client proxy integrisan sa Spring Cloud discovery/load balancing-om."),
            ("Idempotency", "Više istih zahteva proizvodi isti poslovni efekat kao jedan."),
            ("IoC/DI", "Container kreira i ubrizgava zavisnosti umesto da ih klasa sama konstruiše."),
            ("JPA/Hibernate", "ORM specifikacija i njena implementacija za mapiranje objekata i relacione baze."),
            ("JWT", "Kompaktan potpisan token sa claims podacima; payload nije tajna."),
            ("Load balancing", "Izbor jedne od više instanci za konkretan zahtev."),
            ("Micrometer", "Vendor-neutral facade kroz koji aplikacija registruje metrike."),
            ("OAuth2/OIDC", "Autorizacioni okvir i identitetski sloj za tokene i prijavu."),
            ("Optimistic/Pessimistic lock", "Detekcija konflikta verzijom naspram zaključavanja reda pre izmene."),
            ("Outbox", "DB tabela događaja upisana atomski sa domenom i kasnije pouzdano objavljena."),
            ("ProblemDetail", "Standardizovana struktura HTTP greške koju Spring direktno podržava."),
            ("Service discovery", "Dinamičko mapiranje logičkog imena servisa na aktivne instance."),
            ("Span/Trace", "Jedna operacija i kompletno end-to-end putovanje distribuiranog zahteva."),
            ("STOMP/WebSocket", "Messaging protokol i dvosmerni transport za server push."),
            ("Testcontainers", "Test biblioteka koja programatski pokreće pravi dependency u Docker kontejneru."),
        ],
        [2700, 6660],
    )

    document.add_heading("20.4 Završna kontrolna lista", level=2)
    add_bullets(document, [
        "Mogu da nacrtam arhitekturu bez gledanja i objasnim granicu svake baze.",
        "Mogu da ispričam svih 16 koraka transfera i navedem tri nivoa idempotentnosti.",
        "Razlikujem Config/Eureka/Gateway/LoadBalancer/Feign i znam kojim redom učestvuju.",
        "Mogu da objasnim CLOSED/OPEN/HALF_OPEN i tačne pragove iz YAML-a.",
        "Razumem exchange, queue, binding, ack, retry, DLQ i after-commit/outbox razliku.",
        "Mogu da objasnim JWT potpis, issuer, JWK, 401/403 i CUSTOMER/ADMIN pravila.",
        "Znam razliku Actuator/Micrometer/Prometheus/Grafana/Zipkin i šta pokazujem u svakom UI-u.",
        "Mogu sam da pokrenem Compose, demo skriptu, testove i kontrolisan fallback scenario.",
        "Mogu otvoreno da navedem ograničenja: H2, direktni portovi, password grant, outbox gap i in-memory WebSocket broker.",
    ], bullets)

    document.add_heading("Zvanična literatura za dalje učenje", level=1)
    sources = [
        "Spring Boot Reference: https://docs.spring.io/spring-boot/reference/",
        "Spring Cloud Config: https://docs.spring.io/spring-cloud-config/reference/",
        "Spring Cloud Netflix/Eureka: https://docs.spring.io/spring-cloud-netflix/reference/",
        "Spring Cloud Gateway: https://docs.spring.io/spring-cloud-gateway/reference/",
        "Spring Cloud OpenFeign: https://docs.spring.io/spring-cloud-openfeign/reference/",
        "Spring Data JPA: https://docs.spring.io/spring-data/jpa/reference/",
        "Spring AMQP: https://docs.spring.io/spring-amqp/reference/",
        "Spring Security OAuth2 Resource Server: https://docs.spring.io/spring-security/reference/servlet/oauth2/resource-server/",
        "Resilience4j documentation: https://resilience4j.readme.io/docs",
        "Keycloak documentation: https://www.keycloak.org/documentation",
        "Docker Compose documentation: https://docs.docker.com/compose/",
        "Micrometer documentation: https://docs.micrometer.io/micrometer/reference/",
        "Prometheus documentation: https://prometheus.io/docs/",
        "Grafana documentation: https://grafana.com/docs/grafana/latest/",
        "Zipkin documentation: https://zipkin.io/pages/quickstart.html",
        "Testcontainers for Java: https://java.testcontainers.org/",
    ]
    add_bullets(document, sources, bullets)
    add_callout(
        document,
        "Poslednja poruka",
        "Na odbrani ne tvrdi da je projekat produkciona banka. Pokaži da razumeš zašto je rešenje dobro za zadatak, koje probleme već rešava i koji su konkretni sledeći koraci za produkcioni nivo.",
        "positive",
    )

    document.save(GUIDE_PATH)


if __name__ == "__main__":
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    build_report()
    build_guide()
    print(REPORT_PATH)
    print(GUIDE_PATH)
